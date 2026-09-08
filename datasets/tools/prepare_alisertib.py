"""Normalize the supplied Alisertib workbook/JSON into registry import tables.

This is an extraction-preparation step, not the authoritative registry build.
The registry importer reads the generated CSV files under
``datasets/extracted/alisertib``.
"""

from __future__ import annotations

import csv
import json
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl

from registry_common import DATASETS, relative, sha256


PAPER_ID = "paper_PMID22302096"
RAW_WORKBOOK = DATASETS / "raw" / "workbooks" / "Alisertib_MLN8237_Registry_Ready.xlsx"
EXTRACTED = DATASETS / "extracted" / "alisertib"
SUPPLIED_JSON = EXTRACTED / "REGISTRY_IMPORT_DATA.json"


def _write_csv(name: str, rows: list[dict[str, Any]], columns: list[str]) -> None:
    path = EXTRACTED / name
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="raise")
        writer.writeheader()
        writer.writerows({column: row.get(column) for column in columns} for row in rows)


def _slug(value: object) -> str:
    text = str(value).lower().replace("β", "beta")
    return "_".join("".join(character if character.isalnum() else " " for character in text).split())


def _condition_key(context_key: str, treatment: str) -> str:
    return f"{context_key}__{_slug(treatment)}"


def _treatment(condition_1: object, condition_2: object, *, in_vivo: bool = False) -> str:
    first = str(condition_1 or "NONE")
    second = str(condition_2 or "NONE")
    if first.upper() in {"NONE", "VEHICLE", "CV"}:
        return "Vehicle"
    if first == "MLN8237" and second == "CDDP":
        return "MLN8237 30 mg/kg + Cisplatin 2 mg/kg" if in_vivo else "MLN8237 0.5 uM + Cisplatin 2.5 uM"
    if first == "MLN8237":
        return "MLN8237 30 mg/kg" if in_vivo else "MLN8237 0.5 uM"
    if first in {"CDDP", "Cisplatin"}:
        return "Cisplatin 2 mg/kg" if in_vivo else "Cisplatin 2.5 uM"
    raise ValueError(f"Unrecognized treatment: {condition_1!r}, {condition_2!r}")


def _context_key(cell_line: str, *, in_vivo: bool = False) -> str:
    return f"{_slug(cell_line)}_{'xenograft' if in_vivo else 'in_vitro'}"


def _condition_steps(treatment: str, *, in_vivo: bool) -> list[dict[str, Any]]:
    if treatment == "Vehicle":
        return [{
            "perturbation_name": "Vehicle", "dose_value": 0, "dose_unit": None,
            "start_time": 0, "end_time": 21 if in_vivo else 24,
            "time_unit": "d" if in_vivo else "h", "sequence_index": 1,
            "notes": "Vehicle control; formulation, route, and schedule were not encoded in the supplied extraction.",
        }]
    steps: list[dict[str, Any]] = []
    if "MLN8237" in treatment:
        steps.append({
            "perturbation_name": "MLN8237", "dose_value": 30 if in_vivo else 0.5,
            "dose_unit": "mg/kg" if in_vivo else "uM", "start_time": 0,
            "end_time": 21 if in_vivo else 24, "time_unit": "d" if in_vivo else "h",
            "sequence_index": 1,
            "notes": "Oral, daily" if in_vivo else "Drug exposure for 24 h; Figure 3 samples then had 48 h drug-free recovery.",
        })
    if "Cisplatin" in treatment:
        steps.append({
            "perturbation_name": "Cisplatin", "dose_value": 2 if in_vivo else 2.5,
            "dose_unit": "mg/kg" if in_vivo else "uM", "start_time": 0,
            "end_time": 21 if in_vivo else 24, "time_unit": "d" if in_vivo else "h",
            "sequence_index": len(steps) + 1,
            "notes": "Intraperitoneal, twice weekly" if in_vivo else "Drug exposure for 24 h; Figure 3 samples then had 48 h drug-free recovery.",
        })
    return steps


def prepare() -> dict[str, int]:
    if not RAW_WORKBOOK.is_file() or not SUPPLIED_JSON.is_file():
        raise FileNotFoundError("The supplied Alisertib workbook and JSON must be preserved before preparation.")

    workbook = openpyxl.load_workbook(RAW_WORKBOOK, data_only=True)
    supplied = json.loads(SUPPLIED_JSON.read_text(encoding="utf-8-sig"))

    contexts: dict[str, dict[str, Any]] = {}
    conditions: dict[str, dict[str, Any]] = {}
    condition_steps: list[dict[str, Any]] = []
    assays: dict[str, dict[str, Any]] = {}
    observations: list[dict[str, Any]] = []

    def ensure_context(cell_line: str, *, in_vivo: bool = False) -> str:
        key = _context_key(cell_line, in_vivo=in_vivo)
        if key not in contexts:
            contexts[key] = {
                "context_key": key,
                "species": "mouse" if in_vivo else "human",
                "cell_line": cell_line,
                "cell_type": None,
                "tissue": "subcutaneous xenograft tumor" if in_vivo else "esophageal adenocarcinoma cells",
                "disease": "esophageal adenocarcinoma",
                "culture_context": "human-cell xenograft in athymic nude mouse" if in_vivo else "in vitro cell culture",
                "notes": "Human tumor cell line in mouse host." if in_vivo else None,
            }
        return key

    def ensure_condition(cell_line: str, treatment: str, *, in_vivo: bool = False) -> str:
        context_key = ensure_context(cell_line, in_vivo=in_vivo)
        key = _condition_key(context_key, treatment)
        if key not in conditions:
            conditions[key] = {
                "condition_key": key,
                "context_key": context_key,
                "condition_label": treatment,
                "notes": "21-day xenograft regimen." if in_vivo else "Treatment exposure ended at 24 h; observation time is stored separately.",
            }
            for step in _condition_steps(treatment, in_vivo=in_vivo):
                condition_steps.append({
                    "condition_step_key": f"{key}__step_{step['sequence_index']}",
                    "condition_key": key,
                    **step,
                })
        return key

    def ensure_assay(key: str, **row: Any) -> str:
        if key not in assays:
            assays[key] = {"assay_key": key, **row}
        return key

    def add_observation(**row: Any) -> None:
        row.setdefault("replicate_count", None)
        row.setdefault("table", None)
        row.setdefault("lane", None)
        row.setdefault("is_censored", False)
        row.setdefault("censoring_limit", None)
        observations.append(row)

    # Western blots: values are subjective visual annotations from the supplied workbook,
    # not objective ROI-based densitometry. Controls are also extraction-normalized values.
    sheet = workbook["Western Blots"]
    headers = [cell.value for cell in sheet[1]]
    for values in sheet.iter_rows(min_row=2, values_only=True):
        source = dict(zip(headers, values))
        cell_line = str(source["Cell_Line"])
        panel = str(source["Panel"])
        treatment = _treatment(source["Condition_1"], source["Condition_2"])
        condition_key = ensure_condition(cell_line, treatment)
        assay_key = ensure_assay(
            f"wb_fig3_{panel.lower()}_{_slug(cell_line)}",
            assay_type="Western blot", assay_name=f"Figure 3{panel} apoptotic and AURKA markers",
            sample_type="whole-cell lysate", measurement_platform="published immunoblot; supplied visual annotation",
            figure="Figure 3", panel=panel, reported_time=72, reported_time_unit="h",
            replicate_count=None, notes="24 h drug exposure followed by 48 h drug-free recovery.",
        )
        uncertainty = source["Uncertainty"] if source["Uncertainty"] not in (None, 0) else None
        add_observation(
            record_id=f"wb_fig3{panel}_{_slug(cell_line)}_{_slug(source['Protein'])}_{_slug(treatment)}",
            context_key=_context_key(cell_line), condition_key=condition_key, assay_key=assay_key,
            observable=str(source["Protein"]), observable_raw_label=str(source["Protein"]),
            value=source["Value"], value_unit="relative band intensity", time_value=72, time_unit="h",
            statistic="visual normalized signal", uncertainty_type="extraction estimate" if uncertainty is not None else None,
            uncertainty_value=uncertainty, normalization="relative band intensity",
            normalization_reference="vehicle control on supplied annotation scale",
            source_key="workbook", figure="Figure 3", panel=panel, lane=None,
            extraction_method="IMAGE_DERIVED", quality_class="SEMI_QUANTITATIVE",
            notes=(str(source["Notes"] or "") + "; supplied visual estimate, not ROI-based densitometry").strip("; "),
        )

    # qRT-PCR: Figure 3 values came from plots; Figure 6 treatment values were
    # explicitly reported in the Results text with SEM.
    sheet = workbook["qRT-PCR"]
    headers = [cell.value for cell in sheet[1]]
    for values in sheet.iter_rows(min_row=2, values_only=True):
        source = dict(zip(headers, values))
        cell_label = str(source["Cell_Line"])
        in_vivo = "Xenograft" in cell_label
        cell_line = cell_label.replace("_Xenograft", "")
        treatment = _treatment(source["Condition_1"], source["Condition_2"], in_vivo=in_vivo)
        condition_key = ensure_condition(cell_line, treatment, in_vivo=in_vivo)
        if in_vivo:
            panel = "B" if str(source["Gene"]).startswith("PUMA") else "C"
            figure = "Figure 6"
            reported_time, reported_unit = 21, "d"
        else:
            panel = "C" if cell_line == "FLO-1" else "D"
            figure = "Figure 3"
            reported_time, reported_unit = 72, "h"
        assay_key = ensure_assay(
            f"qrtpcr_{_slug(figure)}_{panel.lower()}_{_slug(cell_label)}",
            assay_type="qRT-PCR", assay_name=f"{figure}{panel} PUMA/NOXA expression",
            sample_type="xenograft RNA" if in_vivo else "cell RNA",
            measurement_platform="qRT-PCR", figure=figure, panel=panel,
            reported_time=reported_time, reported_time_unit=reported_unit,
            replicate_count=None, notes="Normalized fold expression.",
        )
        is_control = treatment == "Vehicle"
        is_direct = in_vivo and not is_control
        uncertainty = source["Uncertainty"] if source["Uncertainty"] not in (None, 0) else None
        add_observation(
            record_id=f"qrtpcr_{_slug(figure)}{panel.lower()}_{_slug(cell_label)}_{_slug(source['Gene'])}_{_slug(treatment)}",
            context_key=_context_key(cell_line, in_vivo=in_vivo), condition_key=condition_key, assay_key=assay_key,
            observable=str(source["Gene"]), observable_raw_label=str(source["Gene"]),
            value=source["Value"], value_unit="fold expression", time_value=reported_time,
            time_unit=reported_unit, statistic="mean" if is_direct else "normalized plotted value",
            uncertainty_type="SEM" if is_direct else ("extraction estimate" if uncertainty is not None else None),
            uncertainty_value=uncertainty, normalization="fold expression",
            normalization_reference="vehicle control = 1.0", source_key="workbook",
            figure=figure, panel=panel,
            extraction_method="DIRECT_SOURCE" if is_direct else ("COMPUTED_FROM_SOURCE_VALUES" if is_control else "PLOT_DIGITIZED"),
            quality_class="QUANTITATIVE" if is_direct or is_control else "SEMI_QUANTITATIVE",
            notes=str(source["Notes"] or ""),
        )

    # Corrected Figure 4 Day-21 observations supplied separately as JSON.
    assay_by_stub = {row["stable_id_stub"]: row for row in supplied["assays"]}
    condition_by_stub = {row["stable_id_stub"]: row for row in supplied["conditions"]}
    for source in supplied["observations_tumor_growth"]:
        cell_line = source["cell_line"]
        condition_source = condition_by_stub[source["condition_id_stub"]]
        treatment = _treatment(
            "NONE" if condition_source["treatment_group"] == "Vehicle" else
            ("MLN8237" if condition_source["treatment_group"].startswith("MLN8237") else "CDDP"),
            "CDDP" if "+" in condition_source["treatment_group"] else "NONE",
            in_vivo=True,
        )
        condition_key = ensure_condition(cell_line, treatment, in_vivo=True)
        assay_source = assay_by_stub[source["assay_id_stub"]]
        panel = str(source["panel"])
        assay_key = ensure_assay(
            f"tumor_volume_fig4_{panel.lower()}_{_slug(cell_line)}",
            assay_type="caliper tumor volume", assay_name=f"Figure 4{panel} {cell_line} xenograft efficacy",
            sample_type="xenograft tumor", measurement_platform="caliper tumor volume",
            figure="Figure 4", panel=panel, reported_time=21, reported_time_unit="d",
            replicate_count=None, notes=str(assay_source["measurement_interval"]),
        )
        notes = str(source.get("notes") or "")
        notes = notes.replace(" TGI ~80%.", "").replace(" TGI ~60%.", "")
        notes = notes.rstrip(".; ")
        if source.get("pvalue"):
            notes = f"{notes}; reported {source['pvalue']} (comparator retained as described in supplied extraction)"
        add_observation(
            record_id=f"tumor_fig4{panel.lower()}_{_slug(cell_line)}_day21_{_slug(treatment)}",
            context_key=_context_key(cell_line, in_vivo=True), condition_key=condition_key, assay_key=assay_key,
            observable="tumor_volume", observable_raw_label="tumor volume",
            value=source["value"], value_unit="mm3", time_value=21, time_unit="d", statistic="mean",
            uncertainty_type=source.get("uncertainty_type"), uncertainty_value=source.get("uncertainty"),
            normalization="none", normalization_reference=None, source_key="import_json",
            figure="Figure 4", panel=panel,
            extraction_method=source["extraction_method"], quality_class=source["quality_class"],
            notes=notes,
        )

    # Figure 2 cell-cycle annotations.
    sheet = workbook["Cell Cycle"]
    headers = [cell.value for cell in sheet[1]]
    for values in sheet.iter_rows(min_row=2, values_only=True):
        source = dict(zip(headers, values))
        cell_line = str(source["Cell_Line"])
        treatment = _treatment(source["Condition_1"], source["Condition_2"])
        condition_key = ensure_condition(cell_line, treatment)
        assay_key = ensure_assay(
            f"flow_fig2_{_slug(cell_line)}",
            assay_type="Flow cytometry", assay_name="Figure 2 cell-cycle distribution",
            sample_type="cells", measurement_platform="DNA-content flow cytometry",
            figure="Figure 2", panel=None, reported_time=None, reported_time_unit="h",
            replicate_count=None, notes="Workbook contains visually extracted phase percentages.",
        )
        uncertainty = source["Uncertainty"] if source["Uncertainty"] not in (None, 0) else None
        add_observation(
            record_id=f"cell_cycle_fig2_{_slug(cell_line)}_{int(source['Timepoint_h'])}h_{_slug(source['Phase'])}_{_slug(treatment)}",
            context_key=_context_key(cell_line), condition_key=condition_key, assay_key=assay_key,
            observable=f"cell_cycle_{source['Phase']}_fraction", observable_raw_label=str(source["Phase"]),
            value=source["Percent_Cells"], value_unit="percent", time_value=source["Timepoint_h"], time_unit="h",
            statistic="plotted percentage", uncertainty_type="extraction estimate" if uncertainty is not None else None,
            uncertainty_value=uncertainty, normalization=None, normalization_reference=None,
            source_key="workbook", figure="Figure 2", panel=None,
            extraction_method="PLOT_DIGITIZED", quality_class="SEMI_QUANTITATIVE",
            notes=str(source["Notes"] or ""),
        )

    # Figure 3 clonogenic survival. Combination values are explicitly reported
    # in Results text; supplied single-agent values are retained as plot-derived.
    sheet = workbook["Colony Formation"]
    headers = [cell.value for cell in sheet[1]]
    for values in sheet.iter_rows(min_row=2, values_only=True):
        source = dict(zip(headers, values))
        cell_line = str(source["Cell_Line"])
        treatment = _treatment(source["Condition_1"], source["Condition_2"])
        condition_key = ensure_condition(cell_line, treatment)
        if cell_line == "FLO-1":
            figure, panel = "Figure 3", "A"
        elif cell_line == "OE33":
            figure, panel = "Figure 3", "B"
        else:
            figure, panel = "Supplementary Figure 1", "C"
        assay_key = ensure_assay(
            f"clonogenic_{_slug(figure)}_{panel.lower()}_{_slug(cell_line)}",
            assay_type="Clonogenic survival", assay_name=f"{figure}{panel} colony formation",
            sample_type="cells", measurement_platform="clonogenic survival assay",
            figure=figure, panel=panel, reported_time=None, reported_time_unit=None,
            replicate_count=None, notes="24 h drug exposure followed by colony formation in drug-free medium.",
        )
        is_direct = "+" in treatment
        add_observation(
            record_id=f"colony_{_slug(figure)}{panel.lower()}_{_slug(cell_line)}_{_slug(treatment)}",
            context_key=_context_key(cell_line), condition_key=condition_key, assay_key=assay_key,
            observable="clonogenic_survival", observable_raw_label="survival percent of vehicle control",
            value=source["Survival_Percent_of_CV"], value_unit="percent of vehicle control",
            time_value=None, time_unit=None, statistic="mean", uncertainty_type="SEM",
            uncertainty_value=source["Uncertainty_Percent"], normalization="percent of vehicle control",
            normalization_reference="vehicle control = 100%", source_key="workbook",
            figure=figure, panel=panel,
            extraction_method="DIRECT_SOURCE" if is_direct else "PLOT_DIGITIZED",
            quality_class="QUANTITATIVE" if is_direct else "SEMI_QUANTITATIVE",
            notes=str(source["Notes"] or ""),
        )

    context_columns = ["context_key", "species", "cell_line", "cell_type", "tissue", "disease", "culture_context", "notes"]
    condition_columns = ["condition_key", "context_key", "condition_label", "notes"]
    step_columns = ["condition_step_key", "condition_key", "perturbation_name", "dose_value", "dose_unit", "start_time", "end_time", "time_unit", "sequence_index", "notes"]
    assay_columns = ["assay_key", "assay_type", "assay_name", "sample_type", "measurement_platform", "figure", "panel", "reported_time", "reported_time_unit", "replicate_count", "notes"]
    observation_columns = [
        "record_id", "context_key", "condition_key", "assay_key", "observable", "observable_raw_label",
        "value", "value_unit", "time_value", "time_unit", "statistic", "uncertainty_type",
        "uncertainty_value", "replicate_count", "normalization", "normalization_reference", "source_key",
        "figure", "panel", "table", "lane", "extraction_method", "quality_class", "is_censored",
        "censoring_limit", "notes",
    ]
    _write_csv("contexts.csv", sorted(contexts.values(), key=lambda row: row["context_key"]), context_columns)
    _write_csv("conditions.csv", sorted(conditions.values(), key=lambda row: row["condition_key"]), condition_columns)
    _write_csv("condition_steps.csv", sorted(condition_steps, key=lambda row: row["condition_step_key"]), step_columns)
    _write_csv("assays.csv", sorted(assays.values(), key=lambda row: row["assay_key"]), assay_columns)
    _write_csv("observations.csv", sorted(observations, key=lambda row: row["record_id"]), observation_columns)

    counts = {
        "contexts": len(contexts), "conditions": len(conditions), "condition_steps": len(condition_steps),
        "assays": len(assays), "observations": len(observations),
    }
    metadata = {
        "schema_version": 1,
        "paper_id": PAPER_ID,
        "pmid": "22302096",
        "pmcid": "PMC3297687",
        "doi": "10.1158/1535-7163.MCT-11-0623",
        "extraction_date": "2026-09-07",
        "generated_counts": counts,
        "observation_counts_by_method": dict(sorted(Counter(row["extraction_method"] for row in observations).items())),
        "source_artifacts": [
            {"source_key": "workbook", "path": relative(RAW_WORKBOOK), "sha256": sha256(RAW_WORKBOOK)},
            {"source_key": "import_json", "path": relative(SUPPLIED_JSON), "sha256": sha256(SUPPLIED_JSON)},
        ],
        "caveats": [
            "The original paper PDF is not yet present under datasets/raw/papers.",
            "No intermediate Figure 4 tumor-volume points were fabricated; only supplied Day-21 values are included.",
            "Subjective blot values are IMAGE_DERIVED and SEMI_QUANTITATIVE, not objective densitometry.",
            "The supplied TGI approximations were excluded because no calculation definition was documented.",
        ],
    }
    (EXTRACTED / "extraction_metadata.json").write_text(
        json.dumps(metadata, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    figure4 = {
        "paper_id": PAPER_ID,
        "source_location": "Figure 4 and Results text",
        "panels": {"A": "FLO-1 xenograft", "B": "OE33 xenograft"},
        "x_axis": "time", "x_unit": "day", "y_axis": "tumor volume", "y_unit": "mm3",
        "digitization_method": "Supplied manual curve reading for vehicle endpoints; treatment endpoints transcribed from Results text.",
        "notes": "Only Day-21 observations supplied. Intermediate curve points require separate calibrated plot digitization.",
    }
    (EXTRACTED / "figure_4_provenance.json").write_text(
        json.dumps(figure4, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    return counts


if __name__ == "__main__":
    for table, count in prepare().items():
        print(f"prepared {table}: {count}")
