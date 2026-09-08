"""Normalize the supplied Barasertib extraction into registry-v1 CSV tables.

The supplied workbook, JSON, and draft importer are immutable provenance. This
script corrects the generated layer against the open full text for PMID
27496133/PMCID PMC5050114; it never rewrites the supplied files.
"""

from __future__ import annotations

import csv
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl

from registry_common import DATASETS, relative, sha256


PAPER_ID = "paper_PMID27496133"
EXTRACTED = DATASETS / "extracted" / "barasertib"
RAW_WORKBOOK = DATASETS / "raw" / "workbooks" / "BARASERTIB_OBSERVATIONS.xlsx"
SUPPLIED_JSON = EXTRACTED / "REGISTRY_IMPORT_DATA_BARASERTIB.json"
SUPPLIED_IMPORTER = EXTRACTED / "supplied_barasertib_importer.py"
RAW_PAPER = DATASETS / "raw" / "papers" / "PMID_27496133_PMC5050114.html"
FIGURE_DIR = DATASETS / "raw" / "images" / "barasertib"


AMPLIFICATION = {
    "H82": "MYC", "H211": "MYC", "H446": "MYC", "N417": "MYC",
    "H524": "MYC", "H2171": "MYC",
    "H378": "MYCL", "H748": "MYCL", "H1092": "MYCL", "H1694": "MYCL",
    "H1963": "MYCL", "H2029": "MYCL", "H2141": "MYCL",
    "H69": "MYCN", "H526": "MYCN",
}

SENSITIVE = {"H446", "H211", "N417", "H2171", "H82", "H1963", "H378", "H2081", "H841"}
INTERMEDIATE = {"H524", "H1694", "H1092", "H2029", "H2141", "H69", "H526", "H187", "H146", "DMS114"}
RESISTANT = {"H748", "H345", "H774", "DMS53"}

# Direct values from Table 2 for the subset represented in the supplied JSON.
PLOIDY_CORRECTIONS = {
    "obs_H446_polyploidy_24h_30nm": (30.0, 24.0, 75.0),
    "obs_H446_polyploidy_24h_30nm_8n": (30.0, 24.0, 10.0),
    "obs_H446_polyploidy_48h_30nm": (30.0, 48.0, 11.0),
    "obs_H446_polyploidy_48h_30nm_8n": (30.0, 48.0, 71.0),
    "obs_H69_polyploidy_48h_30nm": (30.0, 48.0, 38.0),
    "obs_H69_polyploidy_48h_30nm_8n": (30.0, 48.0, 58.0),
    "obs_H187_polyploidy_48h_30nm": (30.0, 48.0, 53.0),
    "obs_H187_polyploidy_48h_30nm_8n": (30.0, 48.0, 45.0),
    "obs_H345_polyploidy_48h_30nm": (30.0, 48.0, 72.0),
    "obs_H345_polyploidy_48h_30nm_8n": (30.0, 48.0, 13.0),
    "obs_DMS53_polyploidy_48h_30nm": (50.0, 48.0, 76.0),
    "obs_DMS53_polyploidy_48h_30nm_8n": (50.0, 48.0, 6.0),
}


def _write_csv(name: str, rows: list[dict[str, Any]], columns: list[str]) -> None:
    with (EXTRACTED / name).open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="raise")
        writer.writeheader()
        writer.writerows({column: row.get(column) for column in columns} for row in rows)


def _slug(value: object) -> str:
    return "_".join(re.sub(r"[^a-z0-9]+", " ", str(value).lower()).split())


def _clean(value: object) -> str:
    text = str(value or "")
    for bad, good in {
        "â‰¥": "≥", "â†’": "→", "Ã—": "×", "Â±": "±", "mmÂ³": "mm3",
    }.items():
        text = text.replace(bad, good)
    return text


def _context_key(label: str, *, xenograft: bool = False) -> str:
    return f"{_slug(label)}_{'xenograft' if xenograft else 'in_vitro'}"


def prepare() -> dict[str, int]:
    sources = [RAW_WORKBOOK, SUPPLIED_JSON, SUPPLIED_IMPORTER, RAW_PAPER]
    sources.extend(FIGURE_DIR / f"Helfrich_2016_Figure_{number}.jpg" for number in range(1, 5))
    missing = [str(path) for path in sources if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"Missing Barasertib source artifacts: {missing}")

    workbook = openpyxl.load_workbook(RAW_WORKBOOK, read_only=True, data_only=True)
    if "Barasertib Observations" not in workbook.sheetnames or workbook["Barasertib Observations"].max_row != 43:
        raise ValueError("Unexpected supplied Barasertib workbook layout")
    supplied = json.loads(SUPPLIED_JSON.read_text(encoding="utf-8-sig"))
    if len(supplied.get("observations", [])) != 42:
        raise ValueError("Expected 42 supplied Barasertib observations")

    individual_lines = sorted(
        {row["cell_line"] for row in supplied["observations"] if row["cell_line"] not in {
            "All 23 SCLC lines", "Sensitive lines (n=9)", "Intermediate/Resistant lines (n=14)"
        }} | {"DMS114"}
    )
    contexts: dict[str, dict[str, Any]] = {}
    for cell_line in individual_lines:
        key = _context_key(cell_line)
        amp = AMPLIFICATION.get(cell_line)
        contexts[key] = {
            "context_key": key, "species": "Homo sapiens", "cell_line": cell_line,
            "cell_type": "small-cell lung cancer cell", "tissue": "lung",
            "disease": "small-cell lung cancer", "culture_context": "in vitro cell culture",
            "notes": f"MYC-family amplification: {amp}." if amp else "No MYC-family amplification reported in Table 1.",
        }
    for label, notes in (
        ("SCLC panel n=23", "Aggregate analysis across 23 SCLC cell lines."),
        ("Sensitive cohort n=9", "Aggregate source-defined sensitive cohort."),
        ("Intermediate resistant cohort n=14", "Aggregate intermediate/resistant cohort."),
    ):
        key = _context_key(label)
        contexts[key] = {
            "context_key": key, "species": "Homo sapiens", "cell_line": label,
            "cell_type": "small-cell lung cancer cell cohort", "tissue": "lung",
            "disease": "small-cell lung cancer", "culture_context": "in vitro aggregate cohort",
            "notes": notes,
        }
    xenograft_key = _context_key("H841", xenograft=True)
    contexts[xenograft_key] = {
        "context_key": xenograft_key, "species": "Mus musculus", "cell_line": "H841",
        "cell_type": "human SCLC xenograft", "tissue": "subcutaneous flank tumor",
        "disease": "small-cell lung cancer", "culture_context": "human-cell xenograft in athymic nude mouse",
        "notes": "H841 cells implanted subcutaneously; 5-7 tumor-bearing mice per treatment group.",
    }

    assay_rows = [
        {"assay_key": "mts_growth", "assay_type": "MTS growth assay", "assay_name": "Figure 1 barasertib-HQPA growth inhibition", "sample_type": "cells", "measurement_platform": "CellTiter Aqueous One Solution MTS, absorbance 490 nm", "figure": "Figure 1", "panel": "A-C", "reported_time": 120, "reported_time_unit": "h", "replicate_count": None, "notes": "0-100 nM barasertib-HQPA for five days."},
        {"assay_key": "facs_ploidy", "assay_type": "Flow cytometry", "assay_name": "Table 2 DNA ploidy", "sample_type": "cells", "measurement_platform": "propidium iodide DNA-content flow cytometry", "figure": None, "panel": None, "reported_time": None, "reported_time_unit": "h", "replicate_count": None, "notes": "Direct Table 2 percentages at 24 or 48 h."},
        {"assay_key": "ccle_myc_expression", "assay_type": "Gene expression", "assay_name": "CCLE cMYC expression association", "sample_type": "cell-line expression cohort", "measurement_platform": "CCLE expression data", "figure": "Figure 2", "panel": "A", "reported_time": None, "reported_time_unit": None, "replicate_count": None, "notes": "Source reports cohort means and association statistics."},
        {"assay_key": "myc_amplification_association", "assay_type": "Genomic association", "assay_name": "MYC amplification and sensitivity association", "sample_type": "cell-line cohort", "measurement_platform": "published amplification annotations and Fisher exact test", "figure": "Figure 1", "panel": "A-C", "reported_time": None, "reported_time_unit": None, "replicate_count": 23, "notes": "MYC-family amplification sources are cited in the paper."},
        {"assay_key": "ph3_facs", "assay_type": "Flow cytometry", "assay_name": "Figure 3 phospho-histone H3 Ser10", "sample_type": "cells", "measurement_platform": "intracellular phospho-H3 Ser10 flow cytometry", "figure": "Figure 3", "panel": "A-C", "reported_time": 24, "reported_time_unit": "h", "replicate_count": None, "notes": "Paclitaxel and barasertib-HQPA were used at equal concentrations."},
        {"assay_key": "xenograft_volume", "assay_type": "Caliper tumor volume", "assay_name": "Figure 4 H841 xenograft efficacy", "sample_type": "xenograft tumor", "measurement_platform": "digital caliper", "figure": "Figure 4", "panel": "A", "reported_time": None, "reported_time_unit": "d", "replicate_count": None, "notes": "Tumor volume measured three times weekly."},
    ]

    conditions: dict[str, dict[str, Any]] = {}
    steps: list[dict[str, Any]] = []

    def ensure_condition(
        context_key: str, label: str, step_specs: list[dict[str, Any]], notes: str | None = None
    ) -> str:
        key = f"{context_key}__{_slug(label)}"
        if key not in conditions:
            conditions[key] = {
                "condition_key": key, "context_key": context_key,
                "condition_label": label, "notes": notes,
            }
            for index, spec in enumerate(step_specs, start=1):
                steps.append({
                    "condition_step_key": f"{key}__step_{index}", "condition_key": key,
                    "sequence_index": index, **spec,
                })
        return key

    def analytical_condition(context_key: str) -> str:
        return ensure_condition(context_key, "No experimental perturbation", [{
            "perturbation_name": "No perturbation", "dose_value": 0, "dose_unit": None,
            "start_time": None, "end_time": None, "time_unit": None,
            "notes": "Observational genomic/expression analysis.",
        }])

    observations: list[dict[str, Any]] = []

    def add_observation(**row: Any) -> None:
        defaults = {
            "value": None, "value_unit": None, "time_value": None, "time_unit": None,
            "statistic": None, "uncertainty_type": None, "uncertainty_value": None,
            "replicate_count": None, "normalization": None, "normalization_reference": None,
            "figure": None, "panel": None, "table": None, "lane": None,
            "is_censored": False, "censoring_limit": None, "notes": None,
        }
        defaults.update(row)
        observations.append(defaults)

    # Source-backed growth response classification. The supplied exact IC50
    # guesses are not retained because the paper only reports IC50 <50 nM for
    # sensitive lines and response classes for the other lines.
    response_estimates = {
        source["cell_line"]: source.get("value")
        for source in supplied["observations"] if source["measurement_name"] == "IC50"
    }
    response_estimates["DMS114"] = None
    for cell_line, supplied_estimate in sorted(response_estimates.items()):
        context_key = _context_key(cell_line)
        if cell_line in SENSITIVE:
            condition_key = ensure_condition(context_key, "Barasertib-HQPA 0-100 nM, 120 h", [{
                "perturbation_name": "Barasertib-HQPA", "dose_value": None, "dose_unit": "nM",
                "start_time": 0, "end_time": 120, "time_unit": "h",
                "notes": "Five-day concentration-response experiment.",
            }])
            add_observation(
                record_id=f"{_slug(cell_line)}_barasertib_hqpa_ic50_lt_50_nm",
                context_key=context_key, condition_key=condition_key, assay_key="mts_growth",
                observable="viability_IC50", observable_raw_label="IC50 < 50 nM",
                value_unit="nM", time_value=120, time_unit="h", statistic="IC50",
                source_key="paper", figure="Figure 1", panel="A",
                extraction_method="TEXT_DERIVED", quality_class="QUANTITATIVE",
                is_censored=True, censoring_limit=50,
                notes=f"Source-defined sensitive line: IC50 <50 nM and >75% growth inhibition at 100 nM. Supplied unsupported exact estimate {supplied_estimate} nM was not imported.",
            )
        else:
            phenotype = "intermediate" if cell_line in INTERMEDIATE else "resistant"
            threshold = "32-50% growth inhibition at 100 nM" if phenotype == "intermediate" else "<20% growth inhibition at 100 nM"
            condition_key = ensure_condition(context_key, "Barasertib-HQPA 100 nM, 120 h", [{
                "perturbation_name": "Barasertib-HQPA", "dose_value": 100, "dose_unit": "nM",
                "start_time": 0, "end_time": 120, "time_unit": "h",
                "notes": "Source classification dose and five-day exposure.",
            }])
            add_observation(
                record_id=f"{_slug(cell_line)}_barasertib_hqpa_growth_response_{phenotype}",
                context_key=context_key, condition_key=condition_key, assay_key="mts_growth",
                observable="growth_inhibition_class", observable_raw_label=f"{phenotype}: {threshold}",
                time_value=120, time_unit="h", statistic="source-defined response class",
                source_key="paper", figure="Figure 1", panel="B" if phenotype == "intermediate" else "C",
                extraction_method="TEXT_DERIVED", quality_class="QUALITATIVE_VALIDATION",
                notes=(
                    f"Supplied unsupported exact IC50 estimate {supplied_estimate} nM was not imported."
                    if supplied_estimate is not None
                    else "DMS114 was omitted from the supplied observation rows; its intermediate classification was restored from Figure 1B."
                ),
            )

    # Exact Table 2 ploidy values for the records represented by the input.
    source_by_id = {row["observation_id_stub"]: row for row in supplied["observations"]}
    for record_id, (dose, time, value) in PLOIDY_CORRECTIONS.items():
        source = source_by_id[record_id]
        cell_line = source["cell_line"]
        context_key = _context_key(cell_line)
        condition_key = ensure_condition(context_key, f"Barasertib-HQPA {dose:g} nM, {time:g} h", [{
            "perturbation_name": "Barasertib-HQPA", "dose_value": dose, "dose_unit": "nM",
            "start_time": 0, "end_time": time, "time_unit": "h",
            "notes": "Continuous in vitro exposure before ploidy measurement.",
        }])
        is_8n = "8n" in record_id
        add_observation(
            record_id=record_id, context_key=context_key, condition_key=condition_key,
            assay_key="facs_ploidy", observable="cell_cycle_ge8N_fraction" if is_8n else "cell_cycle_4N_fraction",
            observable_raw_label="DNA content ≥8N" if is_8n else "DNA content 4N",
            value=value, value_unit="percent", time_value=time, time_unit="h",
            statistic="percentage of cells", source_key="paper", table="Table 2",
            extraction_method="DIRECT_SOURCE", quality_class="QUANTITATIVE",
            notes=f"Direct Table 2 value. Supplied phenotype note: {_clean(source.get('phenotype'))}.",
        )

    # Aggregate source-reported MYC statistics.
    panel_context = _context_key("SCLC panel n=23")
    add_observation(
        record_id="myc_amplification_sensitivity_p_value", context_key=panel_context,
        condition_key=analytical_condition(panel_context), assay_key="myc_amplification_association",
        observable="MYC_amplification_sensitivity_p_value", observable_raw_label="MYC amplification association p-value",
        value=0.018, value_unit="dimensionless", statistic="Fisher exact test p-value",
        source_key="paper", figure="Figure 1", panel="A-C", extraction_method="TEXT_DERIVED",
        quality_class="QUANTITATIVE", notes="Odds of sensitivity were 16 times higher (95% CI 1.4-183) for MYC-amplified lines.",
    )
    add_observation(
        record_id="myc_expression_sensitivity_p_value", context_key=panel_context,
        condition_key=analytical_condition(panel_context), assay_key="ccle_myc_expression",
        observable="MYC_expression_sensitivity_p_value", observable_raw_label="cMYC expression association p-value",
        value=0.026, value_unit="dimensionless", statistic="two-group t-test p-value",
        source_key="paper", figure="Figure 2", panel="A", extraction_method="TEXT_DERIVED",
        quality_class="QUANTITATIVE", notes="High expression was associated with sensitivity; source also reports an odds ratio of 11 at a 12.9 cutoff.",
    )
    for label, value, dispersion, context_label in (
        ("sensitive", 10.9, 4.0, "Sensitive cohort n=9"),
        ("intermediate/resistant", 7.2, 3.3, "Intermediate resistant cohort n=14"),
    ):
        context_key = _context_key(context_label)
        add_observation(
            record_id=f"myc_expression_mean_{_slug(label)}", context_key=context_key,
            condition_key=analytical_condition(context_key), assay_key="ccle_myc_expression",
            observable="MYC_mRNA_expression", observable_raw_label=f"mean cMYC expression, {label}",
            value=value, value_unit="CCLE expression score", statistic="mean",
            uncertainty_type="SD", uncertainty_value=dispersion, source_key="paper",
            figure="Figure 2", panel="A", extraction_method="TEXT_DERIVED", quality_class="QUANTITATIVE",
            notes="Results text reports 7.2 for the non-sensitive cohort; the Figure 2 caption reports 7.6, an internal source discrepancy." if label != "sensitive" else None,
        )

    # Correct Figure 3 identities/doses: H446 (25 nM) and H345 (50 nM).
    for cell_line, dose, panel in (("H446", 25.0, "A"), ("H345", 50.0, "B")):
        context_key = _context_key(cell_line)
        condition_key = ensure_condition(context_key, f"Barasertib-HQPA {dose:g} nM + paclitaxel {dose:g} nM, 24 h", [
            {"perturbation_name": "Barasertib-HQPA", "dose_value": dose, "dose_unit": "nM", "start_time": 0, "end_time": 24, "time_unit": "h", "notes": "Concurrent exposure."},
            {"perturbation_name": "Paclitaxel", "dose_value": dose, "dose_unit": "nM", "start_time": 0, "end_time": 24, "time_unit": "h", "notes": "Equal concentration used for mitotic enrichment."},
        ])
        add_observation(
            record_id=f"{_slug(cell_line)}_ph3_ser10_suppressed", context_key=context_key,
            condition_key=condition_key, assay_key="ph3_facs", observable="phospho_H3_Ser10_status",
            observable_raw_label="suppressed relative to paclitaxel alone", time_value=24, time_unit="h",
            statistic="qualitative distribution shift", source_key="figure_3", figure="Figure 3", panel=panel,
            extraction_method="IMAGE_DERIVED", quality_class="QUALITATIVE_VALIDATION",
            notes="Confirms AURKB target engagement in both a sensitive and a resistant line.",
        )

    # Exact tumor volumes from Results plus the source-reported regression status.
    for record_id, label, perturbation, dose, day, value, dispersion in (
        ("h841_vehicle_day34_tumor_volume", "Vehicle, treatment days 20-31", "Vehicle", 0.0, 34.0, 2774.0, 2106.0),
        ("h841_barasertib_50_day34_tumor_volume", "Barasertib 50 mg/kg/day, treatment days 20-31", "Barasertib", 50.0, 34.0, 232.0, 186.0),
        ("h841_barasertib_50_day61_tumor_volume", "Barasertib 50 mg/kg/day, treatment days 20-31", "Barasertib", 50.0, 61.0, 2828.0, 3670.0),
    ):
        condition_key = ensure_condition(xenograft_key, label, [{
            "perturbation_name": perturbation, "dose_value": dose,
            "dose_unit": None if perturbation == "Vehicle" else "mg/kg/day",
            "start_time": 20, "end_time": 31, "time_unit": "d",
            "notes": "Five daily doses, weekend rest, then five daily doses; administration route not reported in supplied extraction.",
        }])
        add_observation(
            record_id=record_id, context_key=xenograft_key, condition_key=condition_key,
            assay_key="xenograft_volume", observable="tumor_volume", observable_raw_label="mean tumor volume",
            value=value, value_unit="mm3", time_value=day, time_unit="d", statistic="mean",
            uncertainty_type="reported dispersion (type unspecified)", uncertainty_value=dispersion,
            source_key="paper", figure="Figure 4", panel="A", extraction_method="TEXT_DERIVED",
            quality_class="QUANTITATIVE", notes="Day measured from tumor-cell implantation.",
        )
    condition_key = ensure_condition(xenograft_key, "Barasertib 100 mg/kg/day, treatment days 20-31", [{
        "perturbation_name": "Barasertib", "dose_value": 100, "dose_unit": "mg/kg/day",
        "start_time": 20, "end_time": 31, "time_unit": "d",
        "notes": "Five daily doses, weekend rest, then five daily doses; administration route not reported in supplied extraction.",
    }])
    add_observation(
        record_id="h841_barasertib_100_day61_regression", context_key=xenograft_key,
        condition_key=condition_key, assay_key="xenograft_volume", observable="tumor_regression_status",
        observable_raw_label="tumors remained regressed through day 61", time_value=61, time_unit="d",
        statistic="qualitative endpoint", source_key="paper", figure="Figure 4", panel="A",
        extraction_method="TEXT_DERIVED", quality_class="QUALITATIVE_VALIDATION",
        notes="No numeric day-61 tumor volume was reported for the 100 mg/kg group.",
    )

    context_columns = ["context_key", "species", "cell_line", "cell_type", "tissue", "disease", "culture_context", "notes"]
    condition_columns = ["condition_key", "context_key", "condition_label", "notes"]
    step_columns = ["condition_step_key", "condition_key", "perturbation_name", "dose_value", "dose_unit", "start_time", "end_time", "time_unit", "sequence_index", "notes"]
    assay_columns = ["assay_key", "assay_type", "assay_name", "sample_type", "measurement_platform", "figure", "panel", "reported_time", "reported_time_unit", "replicate_count", "notes"]
    observation_columns = [
        "record_id", "context_key", "condition_key", "assay_key", "observable", "observable_raw_label",
        "value", "value_unit", "time_value", "time_unit", "statistic", "uncertainty_type",
        "uncertainty_value", "replicate_count", "normalization", "normalization_reference", "source_key",
        "figure", "panel", "table", "lane", "extraction_method", "quality_class", "is_censored",
        "censoring_limit", "notes",
    ]
    _write_csv("contexts.csv", sorted(contexts.values(), key=lambda row: row["context_key"]), context_columns)
    _write_csv("conditions.csv", sorted(conditions.values(), key=lambda row: row["condition_key"]), condition_columns)
    _write_csv("condition_steps.csv", sorted(steps, key=lambda row: row["condition_step_key"]), step_columns)
    _write_csv("assays.csv", sorted(assay_rows, key=lambda row: row["assay_key"]), assay_columns)
    _write_csv("observations.csv", sorted(observations, key=lambda row: row["record_id"]), observation_columns)

    counts = {
        "contexts": len(contexts), "conditions": len(conditions), "condition_steps": len(steps),
        "assays": len(assay_rows), "observations": len(observations),
    }
    artifacts = [
        {"source_key": "workbook", "path": RAW_WORKBOOK},
        {"source_key": "supplied_json", "path": SUPPLIED_JSON},
        {"source_key": "supplied_importer", "path": SUPPLIED_IMPORTER},
        {"source_key": "paper", "path": RAW_PAPER},
    ] + [
        {"source_key": f"figure_{number}", "path": FIGURE_DIR / f"Helfrich_2016_Figure_{number}.jpg"}
        for number in range(1, 5)
    ]
    metadata = {
        "schema_version": 1, "paper_id": PAPER_ID, "pmid": "27496133",
        "pmcid": "PMC5050114", "doi": "10.1158/1535-7163.MCT-16-0298",
        "extraction_date": "2026-09-08", "generated_counts": counts,
        "observation_counts_by_method": dict(sorted(Counter(row["extraction_method"] for row in observations).items())),
        "source_artifacts": [
            {"source_key": item["source_key"], "path": relative(item["path"]), "sha256": sha256(item["path"])}
            for item in artifacts
        ],
        "corrections": [
            "MTS exposure corrected from supplied 72 h to source-reported 120 h (5 days), using 0-100 nM.",
            "Unsupported exact IC50 guesses were replaced by source-reported censoring/classification.",
            "Source response classes corrected: H2171 and H1963 sensitive; H524 intermediate; H748 resistant.",
            "DMS114, the omitted 23rd screened line, was restored as a source-defined intermediate response from Figure 1B.",
            "H446 24 h ploidy corrected to 75% 4N and 10% ≥8N; DMS53 dose corrected to 50 nM.",
            "Figure 3 target-engagement identity corrected from supplied H841 to H446; equal paclitaxel/barasertib-HQPA doses are 25 nM in H446 and 50 nM in H345.",
            "Day-34 and day-61 numeric xenograft volumes from Results were encoded; qualitative text was not placed in numeric value fields.",
        ],
        "caveats": [
            "The original supplied files remain unchanged and may contain unsupported or incorrect values.",
            "The PMC full-text HTML is preserved as the paper artifact; the PDF endpoint did not provide a valid PDF download.",
            "The article reports an internal discrepancy for the intermediate/resistant cMYC mean: 7.2 in Results versus 7.6 in the Figure 2 caption.",
            "Administration route for xenograft dosing is not explicitly encoded because it was not verified from the supplied extraction.",
        ],
    }
    (EXTRACTED / "extraction_metadata.json").write_text(
        json.dumps(metadata, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    return counts


if __name__ == "__main__":
    for table, count in prepare().items():
        print(f"prepared {table}: {count}")
