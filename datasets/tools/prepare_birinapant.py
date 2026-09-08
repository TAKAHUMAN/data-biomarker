"""Normalize the supplied Birinapant extraction into registry-v1 CSV tables.

The BIRINAPANT_OBSERVATIONS.xlsx workbook is a carefully curated source (each
row's Notes typically state it was verified by direct inspection of the cited
figure), similar to the Cilengitide workbook, and is treated as trustworthy
but spot-checked against the open full text for PMID 23403634/PMCID
PMC3618495. The QSP digitization workbook, the supplied JSON, and the
supplied draft importer are immutable provenance and are preserved unchanged;
none of the raw files are modified by this script.
"""

from __future__ import annotations

import csv
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl

from registry_common import DATASETS, relative, sha256


PAPER_ID = "paper_PMID23403634"
EXTRACTED = DATASETS / "extracted" / "birinapant"
RAW_WORKBOOK = DATASETS / "raw" / "workbooks" / "BIRINAPANT_OBSERVATIONS.xlsx"
RAW_QSP_WORKBOOK = DATASETS / "raw" / "workbooks" / "birinapant_qsp_biomarker_digitization.xlsx"
RAW_PAPER = DATASETS / "raw" / "papers" / "PMID_23403634_PMC3618495.html"
RAW_SUPPLEMENT = DATASETS / "raw" / "supplementary" / "birinapant" / "NIHMS446445-supplement-1.docx"
SUPPLIED_JSON = EXTRACTED / "REGISTRY_IMPORT_DATA_BIRINAPANT.json"
SUPPLIED_IMPORTER = EXTRACTED / "supplied_birinapant_importer.py"
FIGURE_DIR = DATASETS / "raw" / "images" / "birinapant"

# Three-tiered sensitivity phenotype reconciled in BIRINAPANT_SMAC_MIMETIC_MOA.md
# (Figure 1 bracket grouping: 1 + 9 + 7 = 17), which supersedes the abstract's
# "twelve of eighteen" figure. Source: Results text, "one of the seventeen".
SINGLE_AGENT_SENSITIVE = {"WM9"}
COMBINATION_SENSITIVE = {
    "WTH202", "WM793B", "WM1366", "WM164", "451Lu",
    "WM1341D", "WM3130", "WM1985", "WM3854",
}
RESISTANT = {"WM1799", "UACC-62", "WM3670", "1205Lu", "WM3918", "C8161", "WM8"}
GENOTYPE = {"WM9": "BRAFV600E", "451Lu": "BRAFV600E", "WM1366": "NRASQ61L"}


def _write_csv(name: str, rows: list[dict[str, Any]], columns: list[str]) -> None:
    EXTRACTED.mkdir(parents=True, exist_ok=True)
    with (EXTRACTED / name).open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="raise")
        writer.writeheader()
        writer.writerows({column: row.get(column) for column in columns} for row in rows)


def _slug(value: object) -> str:
    return "_".join(re.sub(r"[^a-z0-9]+", " ", str(value).lower()).split())


def _sheet_records(workbook: openpyxl.Workbook, name: str) -> list[dict[str, Any]]:
    rows = list(workbook[name].iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(value).strip() if value is not None else "" for value in rows[0]]
    return [
        dict(zip(headers, row))
        for row in rows[1:]
        if any(value not in (None, "") for value in row)
    ]


def _normalized_sheet(workbook: openpyxl.Workbook, sheet: str, output: str) -> int:
    records = _sheet_records(workbook, sheet)
    if not records:
        return 0
    source_headers = list(records[0])
    columns = [_slug(header) for header in source_headers]
    rows = [
        {column: record[source] for column, source in zip(columns, source_headers)}
        for record in records
    ]
    _write_csv(output, rows, columns)
    return len(rows)


def _num(value: object) -> float | None:
    if value in (None, ""):
        return None
    return float(value)


def _tier(cell_line: str) -> str:
    if cell_line in SINGLE_AGENT_SENSITIVE:
        return "Single agent sensitive"
    if cell_line in COMBINATION_SENSITIVE:
        return "Combination sensitive"
    if cell_line in RESISTANT:
        return "Resistant"
    return "Not classified in the 17-line panel"


def prepare() -> dict[str, int]:
    sources = [RAW_WORKBOOK, RAW_QSP_WORKBOOK, RAW_PAPER, RAW_SUPPLEMENT, SUPPLIED_JSON, SUPPLIED_IMPORTER]
    sources.extend(FIGURE_DIR / f"Krepler_2013_Figure_{n}.jpg" for n in range(1, 7))
    missing = [str(p) for p in sources if not p.is_file()]
    if missing:
        raise FileNotFoundError(f"Missing Birinapant source artifacts: {missing}")

    workbook = openpyxl.load_workbook(RAW_WORKBOOK, read_only=True, data_only=True)
    if "Observations" not in workbook.sheetnames or "Metadata" not in workbook.sheetnames:
        raise ValueError(f"Unexpected Birinapant workbook sheets: {workbook.sheetnames}")
    obs_rows = _sheet_records(workbook, "Observations")
    if len(obs_rows) != 170:
        raise ValueError(f"Expected 170 Birinapant observation rows, found {len(obs_rows)}")

    qsp_workbook = openpyxl.load_workbook(RAW_QSP_WORKBOOK, read_only=True, data_only=True)
    required_qsp = {"Recommended input", "Raw densitometry", "Method and caveats"}
    if not required_qsp.issubset(qsp_workbook.sheetnames):
        raise ValueError(f"Unexpected QSP workbook sheets: {qsp_workbook.sheetnames}")

    annotation_counts = {
        "recommended_qsp_inputs": _normalized_sheet(qsp_workbook, "Recommended input", "recommended_qsp_inputs.csv"),
        "raw_densitometry_digitization": _normalized_sheet(qsp_workbook, "Raw densitometry", "raw_densitometry_digitization.csv"),
        "method_and_caveats": _normalized_sheet(qsp_workbook, "Method and caveats", "method_and_caveats.csv"),
    }
    # cIAP1-remaining % + approximate SEM for the xenograft kinetics (Figure 5B),
    # used below to enrich observations.csv with uncertainty the base workbook omits.
    ciap_sem: dict[tuple[str, float], float] = {}
    for row in _sheet_records(qsp_workbook, "Recommended input"):
        model = str(row["model"]).replace("_xenograft", "")
        ciap_sem[(model, float(row["time_h"]))] = float(row["sem_percent_approx"])

    # ---- contexts -----------------------------------------------------
    contexts: dict[str, dict[str, Any]] = {}

    def ensure_context(key: str, **fields: Any) -> str:
        if key not in contexts:
            contexts[key] = {"context_key": key, **fields}
        return key

    def in_vitro_context(cell_line: str) -> str:
        key = f"{_slug(cell_line)}_in_vitro"
        genotype = GENOTYPE.get(cell_line)
        notes = f"Melanoma phenotype tier (Figure 1 bracket grouping): {_tier(cell_line)}."
        if genotype:
            notes += f" Reported driver genotype: {genotype}."
        if cell_line == "451Lu-BR":
            notes = (
                "451Lu subline with acquired BRAF-inhibitor resistance (RAF isoform "
                "switch plus IGF-1R/PI3K pathway upregulation); used in Figure 6C to "
                "test cross-resistance with birinapant+TNF-alpha sensitivity."
            )
        return ensure_context(
            key, species="Homo sapiens", cell_line=cell_line, cell_type="melanoma cell",
            tissue="cutaneous melanoma", disease="melanoma",
            culture_context="in vitro 2D cell culture, DMEM with 5% FBS", notes=notes,
        )

    def xenograft_context(cell_line: str) -> str:
        key = f"{_slug(cell_line)}_xenograft"
        return ensure_context(
            key, species="Mus musculus", cell_line=cell_line,
            cell_type="human melanoma xenograft", tissue="subcutaneous flank tumor",
            disease="melanoma",
            culture_context="human-cell xenograft in nude (immunodeficient) mice",
            notes=f"Melanoma phenotype tier (in vitro): {_tier(cell_line)}. "
                  "Subcutaneous implant; dosing began once tumors were palpable.",
        )

    def spheroid_context(cell_line: str) -> str:
        key = f"{_slug(cell_line)}_spheroid"
        return ensure_context(
            key, species="Homo sapiens", cell_line=cell_line,
            cell_type="melanoma cell (3D spheroid)", tissue="cutaneous melanoma",
            disease="melanoma",
            culture_context="3D spheroid embedded in collagen matrix mimicking stroma",
            notes=f"Melanoma phenotype tier (2D): {_tier(cell_line)}. "
                  "3D validation of the 2D combination-sensitivity phenotype.",
        )

    sensitive_group_key = ensure_context(
        "birinapant_sensitive_lines_pooled", species="Homo sapiens",
        cell_line="Pooled birinapant-sensitive melanoma lines", cell_type="melanoma cell cohort",
        tissue="cutaneous melanoma", disease="melanoma",
        culture_context="in vitro aggregate cohort (individual lines from Supplementary Table S1)",
        notes="Group-mean RIPK1 mRNA scatter (Figure 3D); individual lines composing this "
              "pool are not enumerated in the digitized figure.",
    )
    resistant_group_key = ensure_context(
        "birinapant_resistant_lines_pooled", species="Homo sapiens",
        cell_line="Pooled birinapant-resistant melanoma lines", cell_type="melanoma cell cohort",
        tissue="cutaneous melanoma", disease="melanoma",
        culture_context="in vitro aggregate cohort (individual lines from Supplementary Table S1)",
        notes="Group-mean RIPK1 mRNA scatter (Figure 3D); individual lines composing this "
              "pool are not enumerated in the digitized figure.",
    )

    # ---- assays ---------------------------------------------------------
    assay_rows = [
        {"assay_key": "mts_ic50", "assay_type": "MTS viability assay", "assay_name": "Figure 1 birinapant IC50 (alone and + TNF-alpha)", "sample_type": "melanoma cells", "measurement_platform": "MTS proliferation assay, 72 h, absorbance 490 nm", "figure": "Figure 1", "panel": None, "reported_time": 72, "reported_time_unit": "h", "replicate_count": 3, "notes": "Bar-chart data labels digitized; mean of biological triplicates per figure legend."},
        {"assay_key": "western_apoptosis_signaling", "assay_type": "Western blot", "assay_name": "Figure 2B apoptosis/signaling immunoblot", "sample_type": "melanoma cells", "measurement_platform": "Immunoblotting (cleaved PARP, NF-kB p65, RIP1); 24 h treatment", "figure": "Figure 2", "panel": "B", "reported_time": 24, "reported_time_unit": "h", "replicate_count": None, "notes": "Qualitative band presence/absence and intensity-shift read by direct panel inspection."},
        {"assay_key": "facs_subg1", "assay_type": "Flow cytometry", "assay_name": "Figure 2C sub-G1 apoptotic fraction", "sample_type": "melanoma cells", "measurement_platform": "Propidium iodide DNA-content FACS", "figure": "Figure 2", "panel": "C", "reported_time": 48, "reported_time_unit": "h", "replicate_count": None, "notes": "Bar heights digitized against 0/20/40/60/80/100% gridlines; no printed data labels."},
        {"assay_key": "facs_annexin_v", "assay_type": "Flow cytometry", "assay_name": "Figure 2D Annexin V positivity", "sample_type": "melanoma cells", "measurement_platform": "Annexin V-FITC/PI FACS", "figure": "Figure 2", "panel": "D", "reported_time": 48, "reported_time_unit": "h", "replicate_count": None, "notes": "Bar heights digitized against 0/20/40/60/80/100% gridlines; no printed data labels."},
        {"assay_key": "western_ciap_invitro", "assay_type": "Western blot", "assay_name": "Figure 3A cIAP1/XIAP degradation kinetics (451Lu)", "sample_type": "451Lu melanoma cells", "measurement_platform": "Immunoblotting, dose (0.1-10 uM) x time (1-24 h) matrix", "figure": "Figure 3", "panel": "A", "reported_time": None, "reported_time_unit": "h", "replicate_count": None, "notes": "cIAP1 essentially undetectable at all tested doses/timepoints; XIAP unchanged. See raw_densitometry_digitization.csv for image-derived band intensities."},
        {"assay_key": "mts_rescue", "assay_type": "MTS viability assay", "assay_name": "Figure 3B/3C caspase and RIP1-kinase inhibitor rescue", "sample_type": "melanoma cells", "measurement_platform": "MTS proliferation assay, 72 h; 1 uM birinapant + 1 ng/mL TNF-alpha +/- inhibitor", "figure": "Figure 3", "panel": "B-C", "reported_time": 72, "reported_time_unit": "h", "replicate_count": None, "notes": "Z-VAD-FMK (pan-caspase) and Necrostatin-1 (RIP1 kinase) rescue relative proliferation."},
        {"assay_key": "ripk1_expression", "assay_type": "Gene expression", "assay_name": "Figure 3D RIPK1 mRNA expression", "sample_type": "melanoma cell panel (Supplementary Table S1)", "measurement_platform": "Not specified beyond scatter plot; GUSB-normalized relative expression", "figure": "Figure 3", "panel": "D", "reported_time": None, "reported_time_unit": None, "replicate_count": None, "notes": "Group means visually estimated from scatter plot stratified by birinapant sensitivity; not a strong predictive biomarker on its own."},
        {"assay_key": "spheroid_live_dead", "assay_type": "Confocal microscopy", "assay_name": "Figure 4A 3D spheroid Live/Dead imaging", "sample_type": "melanoma 3D spheroid", "measurement_platform": "Calcein-AM/ethidium homodimer-1 confocal imaging", "figure": "Figure 4", "panel": "A", "reported_time": 72, "reported_time_unit": "h", "replicate_count": None, "notes": "Qualitative red(dead):green(live) shift read by direct image inspection."},
        {"assay_key": "spheroid_alamar_blue", "assay_type": "Alamar Blue viability assay", "assay_name": "Figure 4B 3D spheroid relative viability", "sample_type": "melanoma 3D spheroid", "measurement_platform": "Resazurin reduction fluorescence, 72 h", "figure": "Figure 4", "panel": "B", "reported_time": 72, "reported_time_unit": "h", "replicate_count": 6, "notes": "n=6 technical replicates per line; birinapant 1 uM + TNF-alpha 1 ng/mL."},
        {"assay_key": "xenograft_tumor_volume", "assay_type": "Caliper tumor volume", "assay_name": "Figure 5A xenograft tumor growth", "sample_type": "451Lu/1205Lu xenograft tumor", "measurement_platform": "Digital caliper; V = length x width^2 / 2; fold-change vs first dosing day", "figure": "Figure 5", "panel": "A", "reported_time": None, "reported_time_unit": "d", "replicate_count": 5, "notes": "n=5 mice/group; IP dosing 3x/week 30 mg/kg birinapant or vehicle."},
        {"assay_key": "western_ciap_xenograft", "assay_type": "Western blot", "assay_name": "Figure 5B xenograft cIAP1 target engagement", "sample_type": "451Lu/1205Lu xenograft tumor", "measurement_platform": "Immunoblot densitometry, cIAP1 normalized to vehicle=100%", "figure": "Figure 5", "panel": "B", "reported_time": None, "reported_time_unit": "h", "replicate_count": None, "notes": "Mice dosed twice at a 48 h interval; tumors harvested 3/6/12/24 h after the last dose. SEM enriched from birinapant_qsp_biomarker_digitization.xlsx Recommended input sheet."},
        {"assay_key": "caspase3_ihc_xenograft", "assay_type": "Immunohistochemistry", "assay_name": "Figure 5C xenograft activated caspase-3 IHC", "sample_type": "451Lu/1205Lu xenograft tumor", "measurement_platform": "IHC, activated caspase-3 antibody, paraffin section", "figure": "Figure 5", "panel": "C", "reported_time": 24, "reported_time_unit": "h", "replicate_count": None, "notes": "24 h post-dose; modest increase in both models, more modest in 1205Lu."},
        {"assay_key": "mts_mechanistic", "assay_type": "MTS viability assay", "assay_name": "Figure 6 mechanistic/combination MTS assays", "sample_type": "melanoma cells", "measurement_platform": "MTS proliferation assay, 72 h, absorbance 490 nm or relative viability", "figure": "Figure 6", "panel": "A-D", "reported_time": 72, "reported_time_unit": "h", "replicate_count": None, "notes": "Covers TNF-alpha mAb blocking (6A), birinapant/TNF-alpha dosing-order schedule (6B), BRAFi-resistant subline cross-resistance (6C), and cisplatin combination (6D)."},
    ]

    # ---- conditions -------------------------------------------------
    conditions: dict[str, dict[str, Any]] = {}
    steps: list[dict[str, Any]] = []

    def ensure_condition(context_key: str, label: str, step_specs: list[dict[str, Any]], notes: str | None = None) -> str:
        key = f"{context_key}__{_slug(label)}"
        if key not in conditions:
            conditions[key] = {"condition_key": key, "context_key": context_key, "condition_label": label, "notes": notes}
            for index, spec in enumerate(step_specs, start=1):
                steps.append({"condition_step_key": f"{key}__step_{index}", "condition_key": key, "sequence_index": index, **spec})
        return key

    def vehicle_condition(context_key: str, label: str = "Vehicle control") -> str:
        return ensure_condition(context_key, label, [{
            "perturbation_name": "Vehicle", "dose_value": None, "dose_unit": None,
            "start_time": None, "end_time": None, "time_unit": None, "notes": None,
        }])

    def birinapant_step(dose: float | None, unit: str, start: float | None, end: float | None, time_unit: str, notes: str | None = None) -> dict[str, Any]:
        return {"perturbation_name": "Birinapant", "dose_value": dose, "dose_unit": unit, "start_time": start, "end_time": end, "time_unit": time_unit, "notes": notes}

    def tnf_step(dose: float, unit: str, start: float | None, end: float | None, time_unit: str, notes: str | None = None) -> dict[str, Any]:
        return {"perturbation_name": "TNF-alpha", "dose_value": dose, "dose_unit": unit, "start_time": start, "end_time": end, "time_unit": time_unit, "notes": notes}

    observations: list[dict[str, Any]] = []
    defaults = {
        "value": None, "value_unit": None, "time_value": None, "time_unit": None,
        "statistic": None, "uncertainty_type": None, "uncertainty_value": None,
        "replicate_count": None, "normalization": None, "normalization_reference": None,
        "figure": None, "panel": None, "table": None, "lane": None,
        "is_censored": False, "censoring_limit": None, "notes": None,
    }

    def add(**row: Any) -> None:
        item = dict(defaults)
        item.update(row)
        observations.append(item)

    def parse_figure_panel(text: str | None) -> tuple[str | None, str | None]:
        if not text:
            return None, None
        match = re.match(r"(Figure\s*\d+)\s*([A-Za-z]*)", str(text))
        if not match:
            return text, None
        figure, panel = match.groups()
        return figure, (panel or None)

    for row in obs_rows:
        obs_id = row["Observation ID"]
        cell_line = str(row["Cell Line"]).strip()
        measurement = str(row["Measurement"]).strip()
        dose = row["Dose"]
        dose_unit = row["Dose Unit"]
        value_raw = row["Value"]
        value_unit = row["Value Unit"]
        method = row["Extraction Method"]
        quality = row["Quality Class"]
        figure_panel = row["Figure Panel"]
        phenotype = row["Phenotype"]
        notes = row["Notes"]
        figure, panel = parse_figure_panel(figure_panel)

        def numeric_or_censored() -> tuple[float | None, bool, float | None]:
            text = str(value_raw)
            if text.startswith(">"):
                return None, True, _num(text[1:])
            try:
                return float(value_raw), False, None
            except (TypeError, ValueError):
                return None, False, None

        # 1) Figure 2B western blot: PARP cleavage / NF-kB p65 / RIP1 depletion.
        if measurement.startswith("Cleaved PARP"):
            ctx = in_vitro_context(cell_line)
            combo = "TNF" in measurement
            label = "Birinapant 1000 nM + TNF-alpha 1 ng/mL, 24 h" if combo else "Birinapant 1000 nM, 24 h"
            step_specs = [birinapant_step(1000, "nM", 0, 24, "h")]
            if combo:
                step_specs.append(tnf_step(1, "ng/mL", 0, 24, "h"))
            cond = ensure_condition(ctx, label, step_specs)
            add(record_id=obs_id, context_key=ctx, condition_key=cond, assay_key="western_apoptosis_signaling",
                observable="PARP_cleavage_status", observable_raw_label=f"{value_raw} (89 kDa cleaved PARP fragment)",
                time_value=24, time_unit="h", statistic="blot presence/absence call",
                source_key="figure_2", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                notes=f"{notes} Phenotype tier: {phenotype}.")
            continue

        if measurement.startswith("NF-κB p65 protein level decrease") or measurement.startswith("NF-kB p65 protein level decrease"):
            ctx = in_vitro_context(cell_line)
            cond = ensure_condition(ctx, "Birinapant 1000 nM, 24 h", [birinapant_step(1000, "nM", 0, 24, "h")])
            add(record_id=obs_id, context_key=ctx, condition_key=cond, assay_key="western_apoptosis_signaling",
                observable="NFkB_p65_protein_level_status", observable_raw_label=str(value_raw),
                time_value=24, time_unit="h", statistic="blot intensity call",
                source_key="figure_2", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                notes=f"{notes} Phenotype tier: {phenotype}.")
            continue

        if measurement.startswith("RIP1 kinase protein depletion"):
            ctx = in_vitro_context(cell_line)
            cond = ensure_condition(ctx, "Birinapant 1000 nM + TNF-alpha 1 ng/mL, 24 h", [
                birinapant_step(1000, "nM", 0, 24, "h"), tnf_step(1, "ng/mL", 0, 24, "h"),
            ])
            add(record_id=obs_id, context_key=ctx, condition_key=cond, assay_key="western_apoptosis_signaling",
                observable="RIP1_protein_depletion_status", observable_raw_label=str(value_raw),
                time_value=24, time_unit="h", statistic="blot intensity call",
                source_key="figure_2", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                notes=f"{notes} Phenotype tier: {phenotype}.")
            continue

        # 2) Figure 4A spheroid Live/Dead.
        if measurement.startswith("Live/Dead staining"):
            ctx = spheroid_context(cell_line)
            cond = ensure_condition(ctx, "Birinapant 1000 nM + TNF-alpha 1 ng/mL, 72 h (spheroid)", [
                birinapant_step(1000, "nM", 0, 72, "h"), tnf_step(1, "ng/mL", 0, 72, "h"),
            ])
            add(record_id=obs_id, context_key=ctx, condition_key=cond, assay_key="spheroid_live_dead",
                observable="live_dead_spheroid_status", observable_raw_label=str(value_raw),
                time_value=72, time_unit="h", statistic="qualitative image call",
                source_key="figure_4", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                notes=f"{notes} Phenotype tier: {phenotype}.")
            continue

        # 3) Figure 5C caspase-3 IHC (row applies to both xenograft models).
        if measurement.startswith("Activated caspase-3"):
            for line in ("451Lu", "1205Lu"):
                ctx = xenograft_context(line)
                cond = ensure_condition(ctx, "Birinapant 30 mg/kg IP, 24 h post-dose", [
                    birinapant_step(30, "mg/kg", None, 24, "h", "IP dosing; 24 h post-dose tumor harvest."),
                ])
                add(record_id=f"{obs_id}_{_slug(line)}", context_key=ctx, condition_key=cond, assay_key="caspase3_ihc_xenograft",
                    observable="caspase3_IHC_status", observable_raw_label=str(value_raw),
                    time_value=24, time_unit="h", statistic="qualitative image call",
                    source_key="figure_5", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                    notes=f"{notes} Applies to both 451Lu and 1205Lu xenografts per the source row; 1205Lu increase is described as more modest.")
            continue

        # 4) Figure 1 IC50 (alone / + TNF-alpha).
        if measurement.startswith("IC50"):
            ctx = in_vitro_context(cell_line)
            combo = "TNF" in measurement
            value, censored, limit = numeric_or_censored()
            label = "Birinapant + TNF-alpha 1 ng/mL, dose-response, 72 h" if combo else "Birinapant alone, dose-response, 72 h"
            step_specs = [birinapant_step(None, "nM", 0, 72, "h", "IC50-determining dose-response, approximately 1-1000 nM range.")]
            if combo:
                step_specs.append(tnf_step(1, "ng/mL", 0, 72, "h"))
            cond = ensure_condition(ctx, label, step_specs)
            add(record_id=obs_id, context_key=ctx, condition_key=cond, assay_key="mts_ic50",
                observable="viability_IC50", observable_raw_label=measurement,
                value=value, value_unit=value_unit, time_value=72, time_unit="h", statistic="IC50",
                source_key="figure_1", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                is_censored=censored, censoring_limit=limit,
                notes=f"{notes} Phenotype tier: {phenotype}.")
            continue

        # 5) Figure 2C/2D Sub-G1 and Annexin V, 4 conditions each.
        subg1_match = re.match(r"Sub-G1 apoptotic fraction \((\w[\w -]*)\)", measurement)
        annexin_match = re.match(r"Annexin V positive cells \((\w[\w -]*)\)", measurement)
        if subg1_match or annexin_match:
            arm = (subg1_match or annexin_match).group(1)
            ctx = in_vitro_context(cell_line)
            if arm == "control":
                cond = vehicle_condition(ctx)
                time_value = 48
            elif arm == "birinapant":
                cond = ensure_condition(ctx, "Birinapant 1000 nM, 48 h", [birinapant_step(1000, "nM", 0, 48, "h")])
                time_value = 48
            elif arm == "TNF-alpha":
                cond = ensure_condition(ctx, "TNF-alpha 1 ng/mL, 48 h", [tnf_step(1, "ng/mL", 0, 48, "h")])
                time_value = 48
            else:  # combination
                cond = ensure_condition(ctx, "Birinapant 1000 nM + TNF-alpha 1 ng/mL, 48 h", [
                    birinapant_step(1000, "nM", 0, 48, "h"), tnf_step(1, "ng/mL", 0, 48, "h"),
                ])
                time_value = 48
            assay_key = "facs_subg1" if subg1_match else "facs_annexin_v"
            observable = "subG1_apoptotic_fraction" if subg1_match else "annexinV_positive_fraction"
            add(record_id=obs_id, context_key=ctx, condition_key=cond, assay_key=assay_key,
                observable=observable, observable_raw_label=f"{arm} arm", value=_num(value_raw), value_unit=value_unit,
                time_value=time_value, time_unit="h", statistic="percentage of gated cells",
                uncertainty_type="visual digitization uncertainty (+/-3-5 percentage points)" if method == "PLOT_DIGITIZED" else None,
                source_key="figure_2", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                notes=f"{notes} Phenotype tier: {phenotype}.")
            continue

        # 6) Figure 3A cIAP1 in vitro degradation kinetics (451Lu).
        ciap_kinetics = re.match(r"cIAP1 protein level \((\d+)h, ([\d.]+)uM\)", measurement)
        if ciap_kinetics:
            hours, dose_um = ciap_kinetics.groups()
            ctx = in_vitro_context(cell_line)
            cond = ensure_condition(ctx, f"Birinapant {dose_um} uM, {hours} h", [
                birinapant_step(float(dose_um), "uM", 0, float(hours), "h"),
            ])
            add(record_id=obs_id, context_key=ctx, condition_key=cond, assay_key="western_ciap_invitro",
                observable="cIAP1_protein_level_status", observable_raw_label=str(value_raw),
                time_value=float(hours), time_unit="h", statistic="blot intensity call",
                source_key="figure_3", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                notes=f"{notes} Phenotype tier: {phenotype}.")
            continue

        # 7) Figure 3B/3C rescue curves.
        rescue_match = re.match(r"Relative proliferation \((Z-VAD-FMK|Necrostatin-1) ([\d.]+)uM\)", measurement)
        if rescue_match:
            inhibitor, inhib_dose = rescue_match.groups()
            ctx = in_vitro_context(cell_line)
            label = f"Birinapant 1 uM + TNF-alpha 1 ng/mL + {inhibitor} {inhib_dose} uM, 72 h"
            cond = ensure_condition(ctx, label, [
                birinapant_step(1000, "nM", 0, 72, "h"), tnf_step(1, "ng/mL", 0, 72, "h"),
                {"perturbation_name": inhibitor, "dose_value": float(inhib_dose), "dose_unit": "uM", "start_time": 0, "end_time": 72, "time_unit": "h", "notes": None},
            ])
            add(record_id=obs_id, context_key=ctx, condition_key=cond, assay_key="mts_rescue",
                observable="relative_proliferation_pct", observable_raw_label=measurement,
                value=_num(value_raw), value_unit=value_unit, time_value=72, time_unit="h",
                statistic="percent relative to untreated control",
                source_key="figure_3", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                notes=f"{notes} Phenotype tier: {phenotype}.")
            continue

        # 8) Figure 3D RIPK1 mRNA group means.
        if measurement.startswith("RIPK1 mRNA expression"):
            group = "Sensitive" if "Panel (Sensitive)" == cell_line else "Resistant"
            ctx = sensitive_group_key if group == "Sensitive" else resistant_group_key
            untreated = "untreated" in measurement
            label = "Untreated" if untreated else "TNF-alpha 1 ng/mL"
            step_specs = [{"perturbation_name": "No perturbation", "dose_value": None, "dose_unit": None, "start_time": None, "end_time": None, "time_unit": None, "notes": None}] if untreated else [tnf_step(1, "ng/mL", None, None, None)]
            cond = ensure_condition(ctx, label, step_specs)
            add(record_id=obs_id, context_key=ctx, condition_key=cond, assay_key="ripk1_expression",
                observable="RIPK1_mRNA_relative_expression", observable_raw_label=measurement,
                value=_num(value_raw), value_unit=value_unit, statistic="group mean (visually estimated)",
                source_key="figure_3", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                notes=f"{notes} Phenotype tier: {phenotype}.")
            continue

        # 9) Figure 4B spheroid Alamar Blue.
        if measurement.startswith("Relative viability (Alamar Blue"):
            ctx = spheroid_context(cell_line)
            cond = ensure_condition(ctx, "Birinapant 1000 nM + TNF-alpha 1 ng/mL, 72 h (spheroid)", [
                birinapant_step(1000, "nM", 0, 72, "h"), tnf_step(1, "ng/mL", 0, 72, "h"),
            ])
            add(record_id=obs_id, context_key=ctx, condition_key=cond, assay_key="spheroid_alamar_blue",
                observable="spheroid_relative_viability", observable_raw_label=measurement,
                value=_num(value_raw), value_unit=value_unit, time_value=72, time_unit="h",
                statistic="relative to untreated spheroid control", replicate_count=6,
                source_key="figure_4", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                notes=f"{notes} Phenotype tier: {phenotype}.")
            continue

        # 10) Figure 5A xenograft tumor volume fold-change over time.
        xeno_match = re.match(r"Tumor volume fold-change \(day (\d+), (vehicle|birinapant)\)", measurement)
        if xeno_match:
            day, arm = xeno_match.groups()
            ctx = xenograft_context(cell_line)
            if arm == "vehicle":
                cond = vehicle_condition(ctx, "Vehicle control (xenograft)")
            else:
                cond = ensure_condition(ctx, "Birinapant 30 mg/kg IP 3x/week", [
                    birinapant_step(30, "mg/kg", 0, None, "d", "IP dosing 3x/week (Monday/Wednesday/Friday), started after palpable tumor formation."),
                ])
            add(record_id=obs_id, context_key=ctx, condition_key=cond, assay_key="xenograft_tumor_volume",
                observable="xenograft_tumor_volume_foldchange", observable_raw_label=measurement,
                value=_num(value_raw), value_unit=value_unit, time_value=float(day), time_unit="d",
                statistic="fold change vs day of first dosing",
                source_key="figure_5", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                notes=f"{notes} Phenotype tier (in vitro): {phenotype}.")
            continue

        # 11) Figure 5B xenograft cIAP1 remaining %.
        ciap_xeno_match = re.match(r"cIAP1 protein remaining \((\d+)h post-dose, xenograft\)", measurement)
        if ciap_xeno_match:
            hours = float(ciap_xeno_match.group(1))
            ctx = xenograft_context(cell_line)
            cond = ensure_condition(ctx, "Birinapant 30 mg/kg IP, twice at 48 h interval", [
                birinapant_step(30, "mg/kg", 0, 48, "h", "Dosed twice at a 48 h interval; tumors harvested post final dose."),
            ])
            sem = ciap_sem.get((cell_line, hours))
            add(record_id=obs_id, context_key=ctx, condition_key=cond, assay_key="western_ciap_xenograft",
                observable="cIAP1_protein_remaining_pct", observable_raw_label=measurement,
                value=_num(value_raw), value_unit=value_unit, time_value=hours, time_unit="h",
                statistic="densitometry normalized to vehicle = 100%",
                uncertainty_type="approximate SEM (birinapant_qsp_biomarker_digitization.xlsx Recommended input)" if sem is not None else None,
                uncertainty_value=sem,
                source_key="figure_5", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                notes=f"{notes} Phenotype tier (in vitro): {phenotype}.")
            continue

        # 12) Figure 6A birinapant dose-response absorbance (WM9).
        fig6a_dose = re.match(r"Absorption 490nm \(birinapant dose-response, Fig6A\)", measurement)
        if fig6a_dose:
            ctx = in_vitro_context(cell_line)
            dose_value = _num(dose)
            cond = ensure_condition(ctx, f"Birinapant {dose_value:g} nM, 72 h", [
                birinapant_step(dose_value, "nM", 0, 72, "h"),
            ])
            add(record_id=obs_id, context_key=ctx, condition_key=cond, assay_key="mts_mechanistic",
                observable="MTS_absorbance_490nm", observable_raw_label=measurement,
                value=_num(value_raw), value_unit=value_unit, time_value=72, time_unit="h",
                statistic="absorbance at 490 nm",
                source_key="figure_6", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                notes=f"{notes} Phenotype tier: {phenotype}.")
            continue

        # 13) Figure 6A TNF-alpha mAb blocking (WM9).
        mab_match = re.match(r"Absorption 490nm \(TNF-α mAb ([\d.]+) ug/mL\)", measurement)
        if mab_match:
            mab_dose = float(mab_match.group(1))
            ctx = in_vitro_context(cell_line)
            cond = ensure_condition(ctx, f"Birinapant 1000 nM + TNF-alpha blocking mAb {mab_dose:g} ug/mL, 72 h", [
                birinapant_step(1000, "nM", 0, 72, "h"),
                {"perturbation_name": "TNF-alpha blocking antibody", "dose_value": mab_dose, "dose_unit": "ug/mL", "start_time": 0, "end_time": 72, "time_unit": "h", "notes": "Neutralizing monoclonal antibody against endogenous TNF-alpha."},
            ])
            add(record_id=obs_id, context_key=ctx, condition_key=cond, assay_key="mts_mechanistic",
                observable="MTS_absorbance_490nm_TNF_mAb_blocking", observable_raw_label=measurement,
                value=_num(value_raw), value_unit=value_unit, time_value=72, time_unit="h",
                statistic="absorbance at 490 nm",
                source_key="figure_6", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                notes=f"{notes} Phenotype tier: {phenotype}.")
            continue

        # 14) Figure 6B schedule dependency (451Lu).
        schedule_match = re.match(r"Relative viability \(schedule: (birinapant_then_TNF|TNF_then_birinapant), dose (\d+)nM\)", measurement)
        if schedule_match:
            order, dose_nm = schedule_match.groups()
            dose_nm = float(dose_nm)
            ctx = in_vitro_context(cell_line)
            if order == "birinapant_then_TNF":
                label = f"Schedule: birinapant {dose_nm:g} nM (0-36h) then TNF-alpha 1 ng/mL (36-72h)"
                step_specs = [
                    birinapant_step(dose_nm, "nM", 0, 36, "h", "Dosed first."),
                    tnf_step(1, "ng/mL", 36, 72, "h", "Added after 36 h birinapant pre-incubation."),
                ]
            else:
                label = f"Schedule: TNF-alpha 1 ng/mL (0-36h) then birinapant {dose_nm:g} nM (36-72h)"
                step_specs = [
                    tnf_step(1, "ng/mL", 0, 36, "h", "Dosed first."),
                    birinapant_step(dose_nm, "nM", 36, 72, "h", "Added after 36 h TNF-alpha pre-incubation."),
                ]
            cond = ensure_condition(ctx, label, step_specs)
            add(record_id=obs_id, context_key=ctx, condition_key=cond, assay_key="mts_mechanistic",
                observable="relative_viability_schedule", observable_raw_label=measurement,
                value=_num(value_raw), value_unit=value_unit, time_value=72, time_unit="h",
                statistic="relative to untreated control",
                source_key="figure_6", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                notes=f"{notes} Phenotype tier: {phenotype}.")
            continue

        # 15) Figure 6C 451Lu / 451Lu-BR cross-resistance.
        crossres_match = re.match(r"Relative viability \((451Lu-BR|451Lu) (alone|TNFa), (\d+)nM\)", measurement)
        if crossres_match:
            line_variant, arm, dose_nm = crossres_match.groups()
            dose_nm = float(dose_nm)
            ctx = in_vitro_context(line_variant)
            if arm == "alone":
                cond = ensure_condition(ctx, f"Birinapant {dose_nm:g} nM, 72 h", [birinapant_step(dose_nm, "nM", 0, 72, "h")])
            else:
                cond = ensure_condition(ctx, f"Birinapant {dose_nm:g} nM + TNF-alpha 1 ng/mL, 72 h", [
                    birinapant_step(dose_nm, "nM", 0, 72, "h"), tnf_step(1, "ng/mL", 0, 72, "h"),
                ])
            add(record_id=obs_id, context_key=ctx, condition_key=cond, assay_key="mts_mechanistic",
                observable="relative_viability_braf_crossresistance", observable_raw_label=measurement,
                value=_num(value_raw), value_unit=value_unit, time_value=72, time_unit="h",
                statistic="relative to untreated control",
                source_key="figure_6", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                notes=f"{notes} Phenotype tier: {phenotype}.")
            continue

        # 16) Figure 6D cisplatin combination.
        cisplatin_match = re.match(r"Relative viability \((cisplatin alone|cisplatin birinapant100nM), cisplatin ([\d.]+)uM\)", measurement)
        if cisplatin_match:
            arm, cis_dose = cisplatin_match.groups()
            cis_dose = float(cis_dose)
            ctx = in_vitro_context("451Lu")  # ambiguous "451Lu (or panel line)" cell-line label; see caveats.
            if arm == "cisplatin alone":
                cond = ensure_condition(ctx, f"Cisplatin {cis_dose:g} uM, 72 h", [
                    {"perturbation_name": "Cisplatin", "dose_value": cis_dose, "dose_unit": "uM", "start_time": 0, "end_time": 72, "time_unit": "h", "notes": None},
                ])
            else:
                cond = ensure_condition(ctx, f"Cisplatin {cis_dose:g} uM + birinapant 100 nM, 72 h", [
                    {"perturbation_name": "Cisplatin", "dose_value": cis_dose, "dose_unit": "uM", "start_time": 0, "end_time": 72, "time_unit": "h", "notes": None},
                    birinapant_step(100, "nM", 0, 72, "h"),
                ])
            add(record_id=obs_id, context_key=ctx, condition_key=cond, assay_key="mts_mechanistic",
                observable="relative_viability_cisplatin_combo", observable_raw_label=measurement,
                value=_num(value_raw), value_unit=value_unit, time_value=72, time_unit="h",
                statistic="relative to untreated control",
                source_key="figure_6", figure=figure, panel=panel, extraction_method=method, quality_class=quality,
                notes=f"{notes} Phenotype tier: {phenotype}. Source cell line label is ambiguous ('451Lu (or panel line)'); mapped to the 451Lu in vitro context as the best-supported reading (Figure 6 otherwise centers on 451Lu).")
            continue

        raise ValueError(f"Unclassified Birinapant observation row: {obs_id} / {measurement!r}")

    if len(observations) != 171:  # 170 rows, one (caspase-3 IHC) split across 2 xenograft models
        raise ValueError(f"Expected 171 registry observations (170 rows, 1 split across 2 models), got {len(observations)}")

    # Two additional Figure 5B t=0 vehicle baselines are present in the QSP
    # workbook's Recommended input sheet but omitted from BIRINAPANT_OBSERVATIONS.xlsx.
    for row in _sheet_records(qsp_workbook, "Recommended input"):
        if str(row["treatment"]) != "vehicle":
            continue
        model = str(row["model"]).replace("_xenograft", "")
        ctx = xenograft_context(model)
        cond = vehicle_condition(ctx, "Vehicle control (xenograft)")
        add(record_id=f"qsp_{_slug(model)}_ciap1_vehicle_baseline", context_key=ctx, condition_key=cond,
            assay_key="western_ciap_xenograft", observable="cIAP1_protein_remaining_pct",
            observable_raw_label="cIAP1 protein remaining (0h, vehicle baseline, xenograft)",
            value=float(row["ciap1_percent_vehicle"]), value_unit="% of vehicle control",
            time_value=float(row["time_h"]), time_unit="h",
            statistic="densitometry normalized to vehicle = 100%",
            uncertainty_type="approximate SEM (birinapant_qsp_biomarker_digitization.xlsx Recommended input)",
            uncertainty_value=float(row["sem_percent_approx"]),
            source_key="qsp_workbook", figure="Figure 5", panel="B",
            extraction_method="PLOT_DIGITIZED", quality_class="SEMI_QUANTITATIVE",
            notes="Vehicle t=0 baseline supplied only by the QSP digitization workbook, not by BIRINAPANT_OBSERVATIONS.xlsx; completes the Figure 5B kinetic curve.")

    context_columns = ["context_key", "species", "cell_line", "cell_type", "tissue", "disease", "culture_context", "notes"]
    condition_columns = ["condition_key", "context_key", "condition_label", "notes"]
    step_columns = ["condition_step_key", "condition_key", "perturbation_name", "dose_value", "dose_unit", "start_time", "end_time", "time_unit", "sequence_index", "notes"]
    assay_columns = ["assay_key", "assay_type", "assay_name", "sample_type", "measurement_platform", "figure", "panel", "reported_time", "reported_time_unit", "replicate_count", "notes"]
    observation_columns = [
        "record_id", "context_key", "condition_key", "assay_key", "observable", "observable_raw_label",
        "value", "value_unit", "time_value", "time_unit", "statistic", "uncertainty_type",
        "uncertainty_value", "replicate_count", "normalization", "normalization_reference", "source_key",
        "figure", "panel", "table", "lane", "extraction_method", "quality_class", "is_censored",
        "censoring_limit", "notes",
    ]
    _write_csv("contexts.csv", sorted(contexts.values(), key=lambda r: r["context_key"]), context_columns)
    _write_csv("conditions.csv", sorted(conditions.values(), key=lambda r: r["condition_key"]), condition_columns)
    _write_csv("condition_steps.csv", sorted(steps, key=lambda r: r["condition_step_key"]), step_columns)
    _write_csv("assays.csv", sorted(assay_rows, key=lambda r: r["assay_key"]), assay_columns)
    _write_csv("observations.csv", sorted(observations, key=lambda r: r["record_id"]), observation_columns)

    counts = {
        "contexts": len(contexts), "conditions": len(conditions), "condition_steps": len(steps),
        "assays": len(assay_rows), "observations": len(observations), **annotation_counts,
    }
    artifacts = [
        {"source_key": "workbook", "path": RAW_WORKBOOK},
        {"source_key": "qsp_workbook", "path": RAW_QSP_WORKBOOK},
        {"source_key": "paper", "path": RAW_PAPER},
        {"source_key": "supplement", "path": RAW_SUPPLEMENT},
        {"source_key": "supplied_json", "path": SUPPLIED_JSON},
        {"source_key": "supplied_importer", "path": SUPPLIED_IMPORTER},
    ] + [
        {"source_key": f"figure_{n}", "path": FIGURE_DIR / f"Krepler_2013_Figure_{n}.jpg"} for n in range(1, 7)
    ]
    metadata = {
        "schema_version": 1, "paper_id": PAPER_ID, "pmid": "23403634",
        "pmcid": "PMC3618495", "doi": "10.1158/1078-0432.CCR-12-2518",
        "extraction_date": "2026-09-08", "generated_counts": counts,
        "observation_counts_by_method": dict(sorted(Counter(row["extraction_method"] for row in observations).items())),
        "source_artifacts": [
            {"source_key": item["source_key"], "path": relative(item["path"]), "sha256": sha256(item["path"])}
            for item in artifacts
        ],
        "corrections": [
            "The supplied draft importer (supplied_birinapant_importer.py) carried an incorrect paper title "
            "('Birinapant... kills melanoma cells independently of nitric oxide...'); the correct PMC title "
            "'The novel SMAC mimetic birinapant exhibits potent activity against human melanoma cells' "
            "(matching the workbook Metadata sheet) is used instead.",
            "Spot-checked against PMID 23403634/PMCID PMC3618495 full text: confirmed 'seventeen melanoma cell "
            "lines' and 'effective as a single agent in vitro only in one of the seventeen cell lines tested', "
            "matching the workbook's 1 (single-agent sensitive, WM9) + 9 (combination sensitive) + 7 (resistant) "
            "= 17 tiering used in BIRINAPANT_SMAC_MIMETIC_MOA.md; no numeric corrections to the Observations "
            "workbook were required by this spot check.",
            "Two Figure 5B t=0 vehicle baselines (100% +/- approximate SEM for each of 451Lu and 1205Lu "
            "xenografts), present only in birinapant_qsp_biomarker_digitization.xlsx's Recommended input sheet "
            "and absent from BIRINAPANT_OBSERVATIONS.xlsx, were added to observations.csv to complete the "
            "Figure 5B kinetic curve (record_id prefix qsp_).",
            "Approximate SEM values from the QSP workbook's Recommended input sheet were joined onto the eight "
            "Figure 5B cIAP1-remaining observations (matched by model and timepoint) since the base workbook "
            "does not carry an Uncertainty value for those rows.",
        ],
        "caveats": [
            "IC50 values, sub-G1/Annexin V percentages, cIAP1 densitometry, and most Figure 3/4/5/6 curve points "
            "are PLOT_DIGITIZED (bar-chart or curve digitization); treat as SEMI_QUANTITATIVE/QUANTITATIVE "
            "approximations rather than author-reported exact values, per the workbook's own Quality Class column.",
            "Figure 3A cIAP1/XIAP kinetics rows are recorded as qualitative blot calls (near-complete degradation "
            "at all tested doses/timepoints); image-derived band intensities are preserved separately in "
            "raw_densitometry_digitization.csv but not promoted to numeric observations because the source method "
            "notes flag lane saturation and compressed-image reconstruction as unreliable for precise EC50 fitting.",
            "The Figure 6D cisplatin-combination cell line is ambiguous in the source workbook ('451Lu (or panel "
            "line)'); it is mapped to the 451Lu in vitro context as the best-supported reading, not a confirmed "
            "identity.",
            "Figure 3D RIPK1 mRNA values are pooled group means across unspecified panel subsets (Supplementary "
            "Table S1), not per-cell-line observations; they cannot be joined back to individual cell-line contexts.",
            "No plasma pharmacokinetic/exposure data, receptor occupancy assay, or dose-to-plasma-concentration "
            "mapping is reported in this source; the 30 mg/kg IP xenograft dose and the 0.1-1000 nM/uM in vitro "
            "concentrations are not linked to a PK model without external data.",
            "This package is rich PD/target-engagement/apoptosis-mechanism evidence, but is not a standalone QSP "
            "calibration dataset: exposure-response linkage still requires external PK data.",
        ],
    }
    (EXTRACTED / "extraction_metadata.json").write_text(
        json.dumps(metadata, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    return counts


if __name__ == "__main__":
    for table, count in prepare().items():
        print(f"prepared {table}: {count}")
