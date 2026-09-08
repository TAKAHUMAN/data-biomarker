"""Normalize the supplied AMG 900/Aurora-inhibitor extraction package.

The supplied JSON, workbook, narrative, and Python file are preserved unchanged.
This script creates deterministic CSV intermediates compatible with registry v1.
Known source-backed metadata and value errors are corrected only in the generated
layer and documented in ``extraction_metadata.json``.
"""

from __future__ import annotations

import csv
import json
import re
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl

from registry_common import DATASETS, relative, sha256


PAPER_ID = "paper_PMID29197031"
EXTRACTED = DATASETS / "extracted" / "amg900"
RAW_WORKBOOK = DATASETS / "raw" / "workbooks" / "AURORA_INHIBITORS_OBSERVATIONS.xlsx"
SUPPLIED_JSON = EXTRACTED / "REGISTRY_IMPORT_DATA_AURORA_INHIBITORS.json"
SUPPLIED_MOA = EXTRACTED / "AURORA_INHIBITORS_MOA.md"
SUPPLIED_IMPORTER = EXTRACTED / "supplied_aurora_inhibitors_importer.py"
FIGURE_DIR = DATASETS / "raw" / "images" / "amg900"


# The first correction is explicit in the paper abstract and Figure 1E. The
# remaining corrections are calibrated manual readings of the publicly served
# article figures. They remain PLOT_DIGITIZED/SEMI_QUANTITATIVE.
VALUE_CORRECTIONS: dict[str, float] = {
    "obs_005_AZD1152_93T449_EC50": 74.5,
    "obs_010_AMG900_SW872_polyploidy_baseline": 19.0,
    "obs_011_AMG900_SW872_polyploidy_25nM": 86.0,
    "obs_012_AMG900_SW872_polyploidy_1000nM": 81.0,
    "obs_013_AZD1152_SW872_polyploidy_100nM": 20.0,
    "obs_014_AZD1152_SW872_polyploidy_1000nM": 81.0,
    "obs_015_MK5108_SW872_polyploidy_1000nM": 28.0,
    "obs_016_AMG900_93T449_polyploidy_baseline": 21.0,
    "obs_017_AMG900_93T449_polyploidy_500nM": 56.0,
    "obs_018_AMG900_93T449_polyploidy_1000nM": 57.0,
    "obs_019_AZD1152_93T449_polyploidy_1000nM": 42.0,
    "obs_020_MK5108_93T449_polyploidy_1000nM": 21.0,
    "obs_021_AURKA_mRNA_SW872": 32.0,
    "obs_022_AURKA_mRNA_93T449": 3.2,
    "obs_023_AURKA_mRNA_HCT116": 40.0,
    "obs_024_AURKB_mRNA_SW872": 200.0,
    "obs_025_AURKB_mRNA_93T449": 4.6,
    "obs_026_AURKB_mRNA_HCT116": 470.0,
    "obs_053_AMG900_SW872_polyploidy_1nM": 18.0,
    "obs_054_AMG900_SW872_polyploidy_5nM": 42.0,
    "obs_055_AMG900_SW872_polyploidy_50nM": 86.0,
    "obs_056_AMG900_SW872_polyploidy_250nM": 89.0,
    "obs_057_AZD1152_SW872_polyploidy_25nM": 7.0,
    "obs_058_AZD1152_SW872_polyploidy_250nM": 39.0,
    "obs_059_AZD1152_SW872_polyploidy_500nM": 64.0,
    "obs_060_MK5108_SW872_polyploidy_250nM": 13.0,
    "obs_061_MK5108_SW872_polyploidy_500nM": 15.0,
    "obs_062_AMG900_93T449_polyploidy_25nM": 39.0,
    "obs_063_AMG900_93T449_polyploidy_100nM": 46.0,
    "obs_064_AZD1152_93T449_polyploidy_100nM": 23.0,
    "obs_065_AZD1152_93T449_polyploidy_500nM": 42.0,
    "obs_066_MK5108_93T449_polyploidy_500nM": 15.0,
}


QRT_PCR_CORRECTIONS: dict[str, dict[str, str]] = {
    "obs_021_AURKA_mRNA_SW872": {
        "cell_line": "SW-872", "figure_panel": "Figure 4A",
        "normalization_reference": "differentiated adipocytes = 1",
        "raw_label": "AURKA fold change versus differentiated adipocytes",
    },
    "obs_022_AURKA_mRNA_93T449": {
        "cell_line": "SW-872", "figure_panel": "Figure 4C",
        "normalization_reference": "93T449 = 1",
        "raw_label": "AURKA fold change in SW-872 relative to 93T449",
    },
    "obs_023_AURKA_mRNA_HCT116": {
        "cell_line": "HCT-116", "figure_panel": "Figure 4A",
        "normalization_reference": "differentiated adipocytes = 1",
        "raw_label": "AURKA fold change versus differentiated adipocytes",
    },
    "obs_024_AURKB_mRNA_SW872": {
        "cell_line": "SW-872", "figure_panel": "Figure 4B",
        "normalization_reference": "differentiated adipocytes = 1",
        "raw_label": "AURKB fold change versus differentiated adipocytes",
    },
    "obs_025_AURKB_mRNA_93T449": {
        "cell_line": "SW-872", "figure_panel": "Figure 4C",
        "normalization_reference": "93T449 = 1",
        "raw_label": "AURKB fold change in SW-872 relative to 93T449",
    },
    "obs_026_AURKB_mRNA_HCT116": {
        "cell_line": "HCT-116", "figure_panel": "Figure 4B",
        "normalization_reference": "differentiated adipocytes = 1",
        "raw_label": "AURKB fold change versus differentiated adipocytes",
    },
}


def _write_csv(name: str, rows: list[dict[str, Any]], columns: list[str]) -> None:
    with (EXTRACTED / name).open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="raise")
        writer.writeheader()
        writer.writerows({column: row.get(column) for column in columns} for row in rows)


def _slug(value: object) -> str:
    return "_".join(re.sub(r"[^a-z0-9]+", " ", str(value).lower()).split())


def _context_key(cell_line: str) -> str:
    return f"{_slug(cell_line)}_in_vitro"


def _figure_panel(label: str) -> tuple[str, str | None]:
    match = re.match(r"Figure\s+(\d+)(.*)", label)
    if not match:
        return label, None
    panel = match.group(2).strip().replace(" ", "") or None
    return f"Figure {match.group(1)}", panel


def _observable(measurement: str) -> str:
    return {
        "EC50": "viability_EC50",
        "% Total viable cells": "cell_viability_fraction",
        "% polyploid cells (>4N)": "cell_cycle_polyploid_fraction",
        "AURKA mRNA fold change": "AURKA_mRNA",
        "AURKB mRNA fold change": "AURKB_mRNA",
        "AURKA protein (48 kDa band)": "AURKA_protein",
        "AURKB protein (39 kDa band)": "AURKB_protein",
    }[measurement]


def prepare() -> dict[str, int]:
    required = [RAW_WORKBOOK, SUPPLIED_JSON, SUPPLIED_MOA, SUPPLIED_IMPORTER]
    required.extend(FIGURE_DIR / f"Noronha_2018_Figure_{number}.webp" for number in range(1, 6))
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"Missing AMG 900 source artifacts: {missing}")

    # Validate both supplied structured formats before generating derived files.
    workbook = openpyxl.load_workbook(RAW_WORKBOOK, read_only=True, data_only=True)
    if "Observations" not in workbook.sheetnames or workbook["Observations"].max_row != 67:
        raise ValueError("Unexpected supplied AMG 900 workbook layout")
    supplied = json.loads(SUPPLIED_JSON.read_text(encoding="utf-8-sig"))
    if len(supplied.get("observations", [])) != 66:
        raise ValueError("Expected 66 supplied Aurora-inhibitor observations")

    contexts = {
        "SW-872": {
            "species": "Homo sapiens", "cell_line": "SW-872", "cell_type": None,
            "tissue": "adipose tissue", "disease": "undifferentiated liposarcoma",
            "culture_context": "in vitro cell culture",
            "notes": "Undifferentiated human liposarcoma cell line.",
        },
        "93T449": {
            "species": "Homo sapiens", "cell_line": "93T449", "cell_type": None,
            "tissue": "adipose tissue", "disease": "well-differentiated liposarcoma",
            "culture_context": "in vitro cell culture",
            "notes": "Well-differentiated human liposarcoma cell line.",
        },
        "HCT-116": {
            "species": "Homo sapiens", "cell_line": "HCT-116", "cell_type": None,
            "tissue": "colon", "disease": "colorectal carcinoma",
            "culture_context": "in vitro positive-control cell culture",
            "notes": "Positive-control human colorectal cancer cell line.",
        },
    }
    context_rows = [
        {"context_key": _context_key(cell_line), **row}
        for cell_line, row in contexts.items()
    ]

    assay_rows = [
        {"assay_key": "mtt_viability", "assay_type": "MTT viability", "assay_name": "Figure 1 total viable cell count", "sample_type": "cells", "measurement_platform": "MTT assay", "figure": "Figure 1", "panel": None, "reported_time": 72, "reported_time_unit": "h", "replicate_count": 3, "notes": "Dose-response and EC50 after 72 h."},
        {"assay_key": "facs_polyploidy", "assay_type": "Flow cytometry", "assay_name": "Figures 2-3 polyploidy", "sample_type": "cells", "measurement_platform": "propidium iodide DNA-content flow cytometry", "figure": "Figures 2-3", "panel": None, "reported_time": 72, "reported_time_unit": "h", "replicate_count": 3, "notes": "Percentage of cells with >4N DNA content after 72 h."},
        {"assay_key": "qrtpcr_aurka", "assay_type": "qRT-PCR", "assay_name": "Figure 4 AURKA expression", "sample_type": "cell RNA", "measurement_platform": "TaqMan comparative Ct qRT-PCR", "figure": "Figure 4", "panel": "A,C", "reported_time": None, "reported_time_unit": None, "replicate_count": 3, "notes": "Baseline expression; comparison denominator varies by panel."},
        {"assay_key": "qrtpcr_aurkb", "assay_type": "qRT-PCR", "assay_name": "Figure 4 AURKB expression", "sample_type": "cell RNA", "measurement_platform": "TaqMan comparative Ct qRT-PCR", "figure": "Figure 4", "panel": "B,C", "reported_time": None, "reported_time_unit": None, "replicate_count": 3, "notes": "Baseline expression; comparison denominator varies by panel."},
        {"assay_key": "western_aurka", "assay_type": "Western blot", "assay_name": "Figure 5 AURKA protein", "sample_type": "whole-cell lysate", "measurement_platform": "published immunoblot; visual annotation", "figure": "Figure 5", "panel": "A,C", "reported_time": None, "reported_time_unit": None, "replicate_count": 3, "notes": "Qualitative band-presence annotation; no ROI densitometry."},
        {"assay_key": "western_aurkb", "assay_type": "Western blot", "assay_name": "Figure 5 AURKB protein", "sample_type": "whole-cell lysate", "measurement_platform": "published immunoblot; visual annotation", "figure": "Figure 5", "panel": "B,D", "reported_time": None, "reported_time_unit": None, "replicate_count": 3, "notes": "Qualitative band-presence annotation; no ROI densitometry."},
    ]
    assay_by_stub = {
        "assay_MTT_viability": "mtt_viability",
        "assay_FACS_polyploidy": "facs_polyploidy",
        "assay_qRTPCR_AURKA": "qrtpcr_aurka",
        "assay_qRTPCR_AURKB": "qrtpcr_aurkb",
        "assay_WB_AURKA_protein": "western_aurka",
        "assay_WB_AURKB_protein": "western_aurkb",
    }

    conditions: dict[str, dict[str, Any]] = {}
    condition_steps: dict[str, dict[str, Any]] = {}

    def ensure_condition(cell_line: str, drug: str | None, dose: float | None, measurement: str) -> str:
        context_key = _context_key(cell_line)
        if not drug:
            perturbation = "DMSO" if "polyploid" in measurement else "Untreated"
            key = f"{context_key}__{_slug(perturbation)}"
            label = f"{perturbation} control"
            dose_value, dose_unit, end_time = 0.0, None, 72.0 if perturbation == "DMSO" else None
            notes = "Baseline control." if perturbation == "DMSO" else "Untreated baseline expression."
        elif dose is None:
            perturbation = drug
            key = f"{context_key}__{_slug(drug)}_dose_response_0_1000_nm_72h"
            label = f"{drug}, 0-1000 nM concentration-response"
            dose_value, dose_unit, end_time = None, "nM", 72.0
            notes = "Concentration-response condition used to estimate EC50; individual dose is not applicable."
        else:
            perturbation = drug
            key = f"{context_key}__{_slug(drug)}_{_slug(dose)}_nm_72h"
            label = f"{drug} {dose:g} nM, 72 h"
            dose_value, dose_unit, end_time = dose, "nM", 72.0
            notes = "Single-agent in vitro exposure."
        if key not in conditions:
            conditions[key] = {
                "condition_key": key, "context_key": context_key,
                "condition_label": label, "notes": notes,
            }
            condition_steps[key] = {
                "condition_step_key": f"{key}__step_1", "condition_key": key,
                "perturbation_name": perturbation, "dose_value": dose_value,
                "dose_unit": dose_unit, "start_time": 0.0 if end_time is not None else None,
                "end_time": end_time, "time_unit": "h" if end_time is not None else None,
                "sequence_index": 1, "notes": notes,
            }
        return key

    observations: list[dict[str, Any]] = []
    for original in supplied["observations"]:
        source = dict(original)
        record_id = source["observation_id_stub"]
        measurement = source["measurement_name"]
        if record_id in QRT_PCR_CORRECTIONS:
            correction = QRT_PCR_CORRECTIONS[record_id]
            source["cell_line"] = correction["cell_line"]
            source["figure_panel"] = correction["figure_panel"]
        if record_id in VALUE_CORRECTIONS:
            source["value"] = VALUE_CORRECTIONS[record_id]

        cell_line = source["cell_line"]
        drug = source.get("drug") or None
        dose = float(source["dose"]) if source.get("dose") is not None else None
        condition_key = ensure_condition(cell_line, drug, dose, measurement)
        figure, panel = _figure_panel(source["figure_panel"])
        assay_key = assay_by_stub[source["assay_id_stub"]]
        is_ec50 = measurement == "EC50"
        is_blot = "protein" in measurement
        is_qpcr = "mRNA" in measurement

        if is_blot:
            method, quality = "IMAGE_DERIVED", "QUALITATIVE_VALIDATION"
            uncertainty_type, uncertainty_value = None, None
            raw_label = "visible band" if "present" in record_id else "minimal/faint band"
            normalization = "qualitative comparison to actin loading control"
            normalization_reference = "actin loading control"
        elif is_ec50:
            method, quality = "DIRECT_SOURCE", "QUANTITATIVE"
            uncertainty_type, uncertainty_value = None, None
            raw_label = measurement
            normalization = None
            normalization_reference = None
        else:
            method, quality = "PLOT_DIGITIZED", "SEMI_QUANTITATIVE"
            uncertainty_type = "extraction estimate" if source.get("uncertainty") is not None else None
            uncertainty_value = source.get("uncertainty")
            raw_label = measurement
            normalization = "percent of DMSO control" if measurement == "% Total viable cells" else None
            normalization_reference = "DMSO control = 100%" if measurement == "% Total viable cells" else None
            if is_qpcr:
                correction = QRT_PCR_CORRECTIONS[record_id]
                raw_label = correction["raw_label"]
                normalization = "comparative Ct fold change"
                normalization_reference = correction["normalization_reference"]
                uncertainty_type, uncertainty_value = None, None

        note_parts = [str(source.get("notes") or "").replace("Ã—", "x")]
        if record_id == "obs_005_AZD1152_93T449_EC50":
            note_parts = ["Corrected from supplied 43.4 nM to source-reported 74.5 nM."]
        elif record_id in VALUE_CORRECTIONS:
            note_parts.append("Generated value corrected by manual reading of the preserved public figure.")
        if source.get("pvalue") is not None:
            note_parts.append(f"Supplied annotation records p-value threshold/value {source['pvalue']}; comparator and exact operator require full-text verification.")
        if not is_ec50 and not is_blot:
            note_parts.append("Approximate plot reading; not a raw replicate measurement.")

        observations.append({
            "record_id": record_id,
            "context_key": _context_key(cell_line),
            "condition_key": condition_key,
            "assay_key": assay_key,
            "observable": _observable(measurement),
            "observable_raw_label": raw_label,
            "value": source.get("value"),
            "value_unit": source.get("value_unit"),
            "time_value": 72 if assay_key in {"mtt_viability", "facs_polyploidy"} else None,
            "time_unit": "h" if assay_key in {"mtt_viability", "facs_polyploidy"} else None,
            "statistic": "EC50" if is_ec50 else ("qualitative observation" if is_blot else "approximate plotted mean"),
            "uncertainty_type": uncertainty_type,
            "uncertainty_value": uncertainty_value,
            "replicate_count": 3,
            "normalization": normalization,
            "normalization_reference": normalization_reference,
            "source_key": f"figure_{figure.split()[-1]}",
            "figure": figure,
            "panel": panel,
            "table": None,
            "lane": None,
            "extraction_method": method,
            "quality_class": quality,
            "is_censored": False,
            "censoring_limit": None,
            "notes": "; ".join(part.strip("; ") for part in note_parts if part.strip("; ")),
        })

    context_columns = ["context_key", "species", "cell_line", "cell_type", "tissue", "disease", "culture_context", "notes"]
    condition_columns = ["condition_key", "context_key", "condition_label", "notes"]
    step_columns = ["condition_step_key", "condition_key", "perturbation_name", "dose_value", "dose_unit", "start_time", "end_time", "time_unit", "sequence_index", "notes"]
    assay_columns = ["assay_key", "assay_type", "assay_name", "sample_type", "measurement_platform", "figure", "panel", "reported_time", "reported_time_unit", "replicate_count", "notes"]
    observation_columns = [
        "record_id", "context_key", "condition_key", "assay_key", "observable", "observable_raw_label",
        "value", "value_unit", "time_value", "time_unit", "statistic", "uncertainty_type",
        "uncertainty_value", "replicate_count", "normalization", "normalization_reference", "source_key",
        "figure", "panel", "table", "lane", "extraction_method", "quality_class", "is_censored",
        "censoring_limit", "notes",
    ]
    _write_csv("contexts.csv", sorted(context_rows, key=lambda row: row["context_key"]), context_columns)
    _write_csv("conditions.csv", sorted(conditions.values(), key=lambda row: row["condition_key"]), condition_columns)
    _write_csv("condition_steps.csv", sorted(condition_steps.values(), key=lambda row: row["condition_step_key"]), step_columns)
    _write_csv("assays.csv", sorted(assay_rows, key=lambda row: row["assay_key"]), assay_columns)
    _write_csv("observations.csv", sorted(observations, key=lambda row: row["record_id"]), observation_columns)

    counts = {
        "contexts": len(context_rows), "conditions": len(conditions),
        "condition_steps": len(condition_steps), "assays": len(assay_rows),
        "observations": len(observations),
    }
    artifacts = [
        {"source_key": "workbook", "path": RAW_WORKBOOK},
        {"source_key": "supplied_json", "path": SUPPLIED_JSON},
        {"source_key": "supplied_moa", "path": SUPPLIED_MOA},
        {"source_key": "supplied_importer", "path": SUPPLIED_IMPORTER},
    ] + [
        {"source_key": f"figure_{number}", "path": FIGURE_DIR / f"Noronha_2018_Figure_{number}.webp"}
        for number in range(1, 6)
    ]
    metadata = {
        "schema_version": 1,
        "paper_id": PAPER_ID,
        "pmid": "29197031",
        "doi": "10.1007/s11626-017-0208-4",
        "extraction_date": "2026-09-07",
        "generated_counts": counts,
        "observation_counts_by_method": dict(sorted(Counter(row["extraction_method"] for row in observations).items())),
        "source_artifacts": [
            {"source_key": item["source_key"], "path": relative(item["path"]), "sha256": sha256(item["path"])}
            for item in artifacts
        ],
        "corrections": [
            "Paper metadata corrected from supplied PMID 28657245/year 2017/volume 53/pages 1-14 to PMID 29197031/year 2018/volume 54/pages 71-84.",
            "AZD1152-HQPA EC50 in 93T449 corrected from 43.4 to 74.5 nM from the abstract and Figure 1E.",
            "Figure 4 qRT-PCR comparisons were remapped to their actual panels and denominators.",
            "Selected visibly discordant polyploidy annotations were re-read from preserved public Figures 2 and 3.",
            "Supplied p53 labels were not imported because they are not supported by this extraction package and HCT-116 was mislabeled as mutant.",
            "Qualitative Western-blot calls were reclassified from BLOT_DENSITOMETRY to IMAGE_DERIVED/QUALITATIVE_VALIDATION.",
        ],
        "caveats": [
            "The full paper PDF is not present; public preview figures and authoritative abstract were used for verification.",
            "Dose-response and bar-chart points are approximate plot readings, not raw replicate data.",
            "The supplied narrative and importer are preserved for provenance but are not executed as authoritative code.",
            "This paper contains AMG 900 plus AZD1152-HQPA and MK-5108 comparators; all three are retained.",
        ],
    }
    (EXTRACTED / "extraction_metadata.json").write_text(
        json.dumps(metadata, indent=2, sort_keys=True, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    return counts


if __name__ == "__main__":
    for table, count in prepare().items():
        print(f"prepared {table}: {count}")
